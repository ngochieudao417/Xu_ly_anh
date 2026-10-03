#!/usr/bin/env python3
"""
Generate reports/README_vi.docx (Word version) from README_vi.md content,
formatted as a short project report (title page + headed sections).

    python scripts/generate_readme_docx.py

Re-run any time README_vi.md changes, to keep the two in sync.
"""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from src import config as cfg

SOURCE_MD = PROJECT_ROOT / "README_vi.md"
OUT_DOCX = cfg.REPORTS_DIR / "README_vi.docx"


def add_code_block(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    for i, line in enumerate(text.split("\n")):
        run = p.add_run(("\n" if i else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_body_with_inline_code(doc: Document, text: str) -> None:
    """Render a paragraph, turning `...` spans into monospace runs and **...** into bold."""
    p = doc.add_paragraph()
    token_pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for m in token_pattern.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = p.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif token.startswith("["):
            label = token[1:token.index("]")]
            run = p.add_run(label)
            run.italic = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def build_docx(markdown_text: str) -> Document:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        'Biến động Phát triển Khu vực: Đo lường tốc độ "Bê tông hóa"\n'
        "đại dự án Sân bay Long Thành & Vùng phụ cận (2018–2026)"
    )
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Tài liệu tổng quan dự án — Phase 1: Tiền xử lý dữ liệu và EDA")
    run.italic = True
    run.font.size = Pt(13)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Ngày tạo tài liệu: {date.today().isoformat()}").font.size = Pt(10)

    doc.add_page_break()

    lines = markdown_text.split("\n")
    i = 0
    in_code_block = False
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                add_code_block(doc, "\n".join(code_lines))
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if line.startswith("# "):
            i += 1
            continue  # already used as the document title above
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.strip().startswith("- "):
            add_body_with_inline_code(doc, line.strip()[2:])
            doc.paragraphs[-1].style = doc.styles["List Bullet"]
        elif re.match(r"^\d+\.\s", line.strip()):
            content = re.sub(r"^\d+\.\s", "", line.strip())
            add_body_with_inline_code(doc, content)
            doc.paragraphs[-1].style = doc.styles["List Number"]
        elif line.strip() == "":
            pass
        else:
            add_body_with_inline_code(doc, line.strip())

        i += 1

    return doc


def main():
    if not SOURCE_MD.exists():
        raise FileNotFoundError(f"{SOURCE_MD} not found; nothing to convert.")
    markdown_text = SOURCE_MD.read_text(encoding="utf-8")
    doc = build_docx(markdown_text)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(f"[done] Wrote {OUT_DOCX}")


if __name__ == "__main__":
    main()
