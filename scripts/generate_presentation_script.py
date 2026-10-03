#!/usr/bin/env python3
"""
Sinh kich ban thuyet trinh (script noi) day du tu A-Z cho de tai
"Do luong toc do be tong hoa khu vuc san bay Long Thanh 2018-2026".

Tong hop tu: task.md (Phase 1), 4 file docx trong task/, Scipt W2.js,
va toan bo ket qua thuc nghiem da chay.

Chay:
    python scripts/generate_presentation_script.py
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

BLACK = RGBColor(0, 0, 0)
FONT = "Times New Roman"
OUT = PROJECT_ROOT / "reports" / "Kich_ban_Thuyet_trinh_LongThanh.docx"
FIG = PROJECT_ROOT / "outputs" / "figures"
GEEFIG = FIG / "gee"
MAPS = PROJECT_ROOT / "outputs" / "maps"


def font(run, size=13, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(a), FONT)


def H(doc, text, level=1):
    h = doc.add_heading("", level=level)
    font(h.add_run(text), size={1: 15, 2: 13.5, 3: 13}.get(level, 13), bold=True)
    return h


def P(doc, text, size=13, bold=False, italic=False, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def SAY(doc, text):
    """Loi thoai - phan nguoi thuyet trinh doc len."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(8)
    font(p.add_run(text), size=13)
    return p


def CUE(doc, text):
    """Ghi chu san khau: chieu hinh gi, lam gi."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run("[ " + text + " ]"), size=11.5, italic=True)
    return p


def TBL(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        font(t.rows[0].cells[i].paragraphs[0].add_run(str(h)), size=11.5, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            font(cells[i].paragraphs[0].add_run(str(v)), size=11.5)
    doc.add_paragraph()
    return t


def IMG(doc, path: Path, caption: str, width=5.5):
    if not path.exists():
        P(doc, f"[Thieu hinh: {path.name}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(c.add_run(caption), size=11, italic=True)


def SLIDE(doc, num, title, minutes):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run(f"SLIDE {num} — {title}"), size=13.5, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(6)
    font(p2.add_run(f"Thời lượng gợi ý: {minutes}"), size=11.5, italic=True)


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(13)
    st.font.color.rgb = BLACK

    # ================= TRANG BÌA =================
    P(doc,
      "KỊCH BẢN THUYẾT TRÌNH\n"
      "ĐO LƯỜNG TỐC ĐỘ BÊ TÔNG HÓA KHU VỰC SÂN BAY LONG THÀNH\n"
      "GIAI ĐOẠN 2018 – 2026",
      size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    P(doc, "Học phần: Xử lý ảnh — Nhóm 07", size=13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    P(doc, f"Ngày soạn: {date.today().strftime('%d/%m/%Y')}", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()

    H(doc, "CÁCH DÙNG TÀI LIỆU NÀY", 1)
    P(doc, "Tài liệu gồm 3 phần:")
    P(doc, "• Phần A — Kịch bản nói theo từng slide. Chữ thường là lời thoại đọc thẳng được; "
            "phần trong ngoặc vuông in nghiêng là ghi chú sân khấu (chiếu hình gì, chỉ vào đâu).")
    P(doc, "• Phần B — Bảng số liệu tra nhanh, để sẵn khi bị hỏi con số cụ thể.")
    P(doc, "• Phần C — Dự kiến câu hỏi phản biện và cách trả lời.")
    P(doc, "Tổng thời lượng thiết kế khoảng 15–18 phút nói, chưa tính hỏi đáp. Nếu bị cắt thời gian "
            "còn 10 phút thì bỏ Slide 5, 11 và 12, giữ nguyên phần còn lại.")

    doc.add_page_break()

    # ================= PHẦN A =================
    H(doc, "PHẦN A — KỊCH BẢN NÓI THEO SLIDE", 1)

    # --- SLIDE 1 ---
    SLIDE(doc, 1, "Mở đầu, giới thiệu đề tài", "1 phút")
    CUE(doc, "Slide tiêu đề: tên đề tài, tên nhóm, danh sách thành viên")
    SAY(doc,
        "Xin chào thầy và các bạn. Nhóm 07 xin trình bày đề tài: Đo lường tốc độ bê tông hóa khu vực "
        "sân bay Long Thành và vùng phụ cận, giai đoạn 2018 đến 2026.")
    SAY(doc,
        "Nói ngắn gọn thì bài toán của nhóm là thế này: dùng ảnh vệ tinh chụp cùng một khu đất qua "
        "nhiều năm, để máy tự phân loại xem chỗ nào là cây cối, chỗ nào là đất trống, chỗ nào là mặt "
        "nước, chỗ nào đã bị bê tông hóa. Sau đó đo xem diện tích bê tông tăng bao nhiêu qua từng mốc "
        "thời gian, và tăng nhanh nhất ở giai đoạn nào.")

    # --- SLIDE 2 ---
    SLIDE(doc, 2, "Lý do chọn đề tài", "1,5 phút")
    CUE(doc, "Slide: 3 gạch đầu dòng lý do + 1 ảnh sân bay Long Thành từ báo chí")
    SAY(doc,
        "Vì sao nhóm chọn khu vực này chứ không phải chỗ khác? Có ba lý do.")
    SAY(doc,
        "Thứ nhất, đây là công trường xây dựng lớn nhất Việt Nam hiện nay. Sân bay Long Thành được "
        "phê duyệt đầu tư giai đoạn 1 theo Quyết định 1777 năm 2020, khởi công đầu năm 2021 và khai "
        "trương tháng 12 năm 2025. Nghĩa là toàn bộ quá trình từ đất nông nghiệp thành sân bay diễn ra "
        "gọn trong đúng khoảng thời gian mà ảnh Sentinel-2 có sẵn dữ liệu. Rất hiếm khi có một trường "
        "hợp nghiên cứu sạch và rõ ràng như vậy.")
    SAY(doc,
        "Thứ hai, đây là bài toán đo được. Bê tông hóa không phải khái niệm mơ hồ, nó là bề mặt không "
        "thấm nước, và bề mặt không thấm nước có đặc trưng phổ khác hẳn cây cối hay đất trống. Nghĩa là "
        "máy phân biệt được, và mình kiểm chứng được bằng số.")
    SAY(doc,
        "Thứ ba, kết quả có ý nghĩa thực tế. Tốc độ đô thị hóa quanh một sân bay lớn là thứ mà quy "
        "hoạch đô thị và quản lý đất đai đều cần biết. Cách làm này rẻ hơn nhiều so với đi khảo sát "
        "thực địa, và làm lại được cho bất kỳ khu vực nào khác.")

    # --- SLIDE 3 ---
    SLIDE(doc, 3, "Mục tiêu và câu hỏi nghiên cứu", "1 phút")
    CUE(doc, "Slide: 4 câu hỏi nghiên cứu dạng gạch đầu dòng")
    SAY(doc,
        "Đề tài đặt ra bốn câu hỏi cần trả lời bằng số liệu.")
    SAY(doc,
        "Một, đến năm 2026 thì bao nhiêu phần trăm diện tích vùng nghiên cứu đã bị bê tông hóa.")
    SAY(doc,
        "Hai, tốc độ bê tông hóa nhanh nhất rơi vào giai đoạn nào.")
    SAY(doc,
        "Ba, mức độ bê tông hóa có phụ thuộc vào khoảng cách tới sân bay hay không, càng gần thì càng "
        "dày đặc hay là lan đều.")
    SAY(doc,
        "Bốn, và quan trọng về mặt phương pháp: kết quả này có đáng tin không, hay chỉ là do một nguồn "
        "dữ liệu duy nhất tạo ra.")

    # --- SLIDE 4 ---
    SLIDE(doc, 4, "Vùng nghiên cứu", "1,5 phút")
    CUE(doc, "Chiếu ảnh RGB toàn AOI (aoi_rgb_2022.png), chỉ tay vào vị trí sân bay ở giữa")
    IMG(doc, GEEFIG / "aoi_rgb_2022.png", "Ảnh tổ hợp màu tự nhiên toàn vùng nghiên cứu, composite mùa khô 2022.", 5.0)
    SAY(doc,
        "Đây là vùng nghiên cứu của nhóm. Một hình chữ nhật 11,87 nhân 8,64 ki lô mét, diện tích khoảng "
        "102 ki lô mét vuông, bao quanh dự án sân bay Long Thành thuộc huyện Long Thành, tỉnh Đồng Nai. "
        "Tâm dự án ở tọa độ 107,041 độ Đông và 10,784 độ Bắc.")
    SAY(doc,
        "Có một chi tiết về phương pháp mà nhóm muốn nhấn mạnh ở đây. Vùng này được cố định bằng tọa độ "
        "tuyệt đối, viết cứng trong code, chứ không phải khoanh vẽ bằng chuột trên bản đồ. Lý do là vì "
        "mình phải so sánh năm mốc thời gian với nhau. Nếu mỗi lần chạy mà ranh giới xê dịch đi vài chục "
        "mét thì phần diện tích chênh lệch đó sẽ bị tính nhầm thành biến động, trong khi thực tế nó chỉ "
        "là lỗi khoanh vùng. Cố định tọa độ tuyệt đối bảo đảm năm bản đồ khớp khít nhau tới từng điểm ảnh.")
    SAY(doc,
        "Ở độ phân giải 10 mét thì vùng này tương ứng ma trận ảnh 1.190 cột nhân 870 hàng.")

    # --- SLIDE 5 ---
    SLIDE(doc, 5, "Dữ liệu sử dụng", "1,5 phút")
    CUE(doc, "Slide: bảng 3 nguồn dữ liệu")
    TBL(doc, ["Nguồn", "Mã bộ dữ liệu", "Độ phân giải", "Vai trò"], [
        ["Sentinel-2 Level-2A", "COPERNICUS/S2_SR_HARMONIZED", "10 m", "Nguồn chính cho toàn bộ phân tích"],
        ["Landsat 8 Collection 2 L2", "LANDSAT/LC08/C02/T1_L2", "30 m", "Kiểm chứng chéo độc lập"],
        ["ESA WorldCover v200", "ESA/WorldCover/v200", "10 m", "Nguồn nhãn huấn luyện (nền 2021)"],
    ])
    SAY(doc,
        "Nhóm dùng ba nguồn dữ liệu, tất cả đều miễn phí và mở.")
    SAY(doc,
        "Nguồn chính là Sentinel-2 ở mức Level-2A, tức là ảnh đã được hiệu chỉnh khí quyển sẵn, độ phân "
        "giải 10 mét. Chỗ này nhóm từng mắc lỗi: ban đầu có thành viên dùng nhầm bộ Level-1C chưa hiệu "
        "chỉnh, dẫn đến giá trị phổ giữa các thành viên không so sánh được với nhau. Sau đó cả nhóm "
        "thống nhất chỉ dùng một bộ duy nhất.")
    SAY(doc,
        "Nguồn thứ hai là Landsat 8, độ phân giải thô hơn, 30 mét, dùng để kiểm chứng chéo. Ý tưởng là "
        "nếu kết quả đúng thì hai vệ tinh khác nhau phải cho ra cùng một xu hướng.")
    SAY(doc,
        "Nguồn thứ ba là ESA WorldCover, một bản đồ lớp phủ toàn cầu, dùng làm nhãn huấn luyện cho mô "
        "hình. Nhờ nó nhóm không phải ngồi khoanh tay hàng nghìn điểm mẫu.")

    # --- SLIDE 6 ---
    SLIDE(doc, 6, "Quy trình tổng thể", "2 phút")
    CUE(doc, "Slide sơ đồ khối: 7 bước từ ảnh thô tới kết quả — đây là slide xương sống, nói chậm")
    P(doc, "Sơ đồ để vẽ lên slide:", italic=True)
    P(doc, "Ảnh thô Sentinel-2 → Ghép ảnh mùa khô + che mây → Bộ 9 đặc trưng → "
           "Lấy mẫu WorldCover + làm sạch nhãn → Huấn luyện Random Forest (A vs B) → "
           "Phân loại 5 mốc + lọc nhiễu → Tính diện tích, vành đai, kiểm chứng Landsat")
    SAY(doc,
        "Đây là toàn bộ quy trình, gồm bảy bước. Em sẽ đi qua từng bước và nói rõ tại sao lại làm như vậy.")
    SAY(doc,
        "Bước một, ghép ảnh. Một ảnh vệ tinh chụp một ngày cụ thể thì rất dễ dính mây, mà vùng Nam Bộ "
        "thì mùa mưa mây gần như quanh năm. Nên nhóm không dùng ảnh một ngày, mà gom tất cả ảnh chụp "
        "trong khung mùa khô, từ tháng 12 năm trước đến hết tháng 4 năm sau, che mây từng ảnh, rồi lấy "
        "giá trị trung vị. Trung vị chứ không phải trung bình, vì trung vị không bị kéo lệch bởi vài "
        "điểm ảnh mây sót lại.")
    SAY(doc,
        "Việc che mây làm ở mức từng điểm ảnh, dựa vào lớp SCL mà chính ESA đã gán sẵn cho mỗi ảnh. "
        "SCL đánh dấu điểm ảnh nào là bóng mây, mây trung bình, mây dày, mây ti mỏng. Nhóm loại bỏ đúng "
        "bốn loại đó.")
    SAY(doc,
        "Bước hai, tạo đặc trưng. Ngoài sáu kênh phổ gốc, nhóm tính thêm ba chỉ số: NDVI cho thảm thực "
        "vật, NDBI cho bề mặt xây dựng, NDWI cho mặt nước. Ba chỉ số này đều là công thức chuẩn trong "
        "viễn thám, lấy hiệu chia tổng của hai kênh phổ. Tổng cộng có chín đặc trưng cho mỗi điểm ảnh.")
    SAY(doc,
        "Ba bước còn lại là lấy mẫu, huấn luyện và áp mô hình, em xin trình bày kỹ ở các slide sau.")

    # --- SLIDE 6B ---
    SLIDE(doc, "6B", "Pipeline chi tiết 11 bước (theo đúng code nhóm chạy)", "2 phút")
    CUE(doc, "Slide: sơ đồ 11 khối, đánh số — chiếu song song với file Scipt W2.js nếu thầy muốn xem code")
    TBL(doc, ["Bước", "Nội dung", "Đầu ra"], [
        ["1", "Khai báo AOI + hằng số dùng chung (tọa độ tuyệt đối, bảng màu, seed=42, ngưỡng mây 60%)", "Vùng nghiên cứu cố định"],
        ["2", "Ghép ảnh Sentinel-2 khung mùa khô: lọc cảnh → mask mây SCL (mã 3/8/9/10) → chia 10.000 → median → clip AOI", "5 ảnh composite"],
        ["3", "Tạo bộ 9 đặc trưng: 6 kênh phổ gốc + NDVI + NDBI + NDWI", "Ảnh 9 kênh mỗi mốc"],
        ["4", "Lấy mẫu từ ESA WorldCover: remap 11 lớp → 4 lớp; 300 điểm/lớp; lọc riêng lớp Water bằng NDWI > 0", "1.200 điểm train/test"],
        ["5", "Huấn luyện Random Forest 200 cây, 2 phương án A (9 band) và B (5 band), so sánh OA/Kappa", "Ma trận nhầm lẫn A và B"],
        ["6", "Tính độ chính xác có trọng số theo diện tích (Olofsson và cộng sự, 2014)", "Con số dùng trong báo cáo"],
        ["7", "Áp mô hình thắng lên 5 mốc + lọc majority 3×3 khử nhiễu muối tiêu", "5 bản đồ phân loại"],
        ["8", "Tính diện tích từng lớp qua các mốc bằng ee.Image.pixelArea()", "Bảng diện tích (ha)"],
        ["9", "Phân tích vành đai khoảng cách 0–2, 2–4, 4–6 km + định lượng phần vành bị cắt", "Tỷ lệ bê tông theo vành"],
        ["10", "Kiểm chứng chéo Landsat 8: hiệu chỉnh Collection 2, mask QA_PIXEL, hạ S2 về 30 m, tính tương quan Pearson", "Hệ số r 4 năm + bản đồ hiệu số"],
        ["11", "Xuất kết quả ra Google Drive (5 ảnh phân loại + ranh giới AOI dạng GeoJSON)", "Sản phẩm bàn giao"],
    ])
    SAY(doc,
        "Slide này là pipeline đầy đủ theo đúng code nhóm chạy, gồm 11 bước, viết trong một file duy "
        "nhất chạy được từ trên xuống dưới. Em không đọc hết 11 dòng, chỉ xin nhấn ba điểm.")
    SAY(doc,
        "Điểm thứ nhất là bước 6, tính độ chính xác có trọng số theo diện tích. Bước này nhóm thêm vào "
        "sau khi tự nhận ra tập kiểm tra cân bằng lớp không phản ánh đúng thực tế, và làm theo khuyến "
        "nghị của Olofsson và cộng sự năm 2014.")
    SAY(doc,
        "Điểm thứ hai là bước 9, trong hàm tính vành đai nhóm có tính luôn tỷ lệ phần vành bị cắt so "
        "với vòng tròn lý thuyết, để nêu minh bạch trong báo cáo chứ không nói chung chung.")
    SAY(doc,
        "Điểm thứ ba là toàn bộ 11 bước đều nằm trong một script duy nhất, cố định seed, nên bất kỳ ai "
        "chạy lại cũng ra đúng con số như trong báo cáo này.")

    # --- SLIDE 6C ---
    SLIDE(doc, "6C", "Cách nhóm chia việc: 9 đầu việc tuần 2", "1,5 phút")
    CUE(doc, "Slide: bảng 9 task — nói lướt, chỉ dừng ở task 3.4 và 3.9")
    TBL(doc, ["Task", "Nội dung", "Kết quả thực tế"], [
        ["3.1", "Chuyển pipeline từ AWS/STAC sang GEE, chạy lại trên AOI chuẩn", "Script GEE thống nhất cho cả nhóm"],
        ["3.2", "Tạo bộ đặc trưng 9 band", "Ảnh 9 kênh cho 5 mốc"],
        ["3.3", "Lấy mẫu WorldCover, 300 điểm/lớp, seed 42, chia 70/30", "1.200 điểm; train 847 / test 353"],
        ["3.4", "Rà mẫu bằng mắt ~20 điểm/lớp", "Phát hiện lớp Water sai ~30%"],
        ["3.5", "Train Random Forest phương án A (9 band)", "OA 85,27% — Kappa 0,803"],
        ["3.6", "Train phương án B (5 band rút gọn)", "OA 81,02% — Kappa 0,746"],
        ["3.7", "So sánh A và B, chọn phương án thắng", "Chọn A trên mọi chỉ số"],
        ["3.8", "Áp mô hình lên 5 mốc, tính diện tích", "5 bản đồ + bảng diện tích"],
        ["3.9", "Kiểm chứng riêng mốc 2018", "Phát hiện và xử lý bất thường 16,2%"],
    ])
    SAY(doc,
        "Đây là cách nhóm chia việc trong tuần thứ hai, chín đầu việc đánh số từ 3.1 đến 3.9. Em xin "
        "dừng ở hai đầu việc mà nhóm thấy có giá trị nhất.")
    SAY(doc,
        "Task 3.4 là rà mẫu bằng mắt. Nhóm đã định bỏ qua bước này vì nghĩ WorldCover là dữ liệu chính "
        "thống của ESA thì chắc đúng. Nhưng khi làm thì phát hiện lớp Water sai tới 30 phần trăm. Nếu "
        "bỏ qua thì toàn bộ kết quả phía sau đã bị lỗi mà không ai biết.")
    SAY(doc,
        "Task 3.9 là kiểm chứng riêng mốc 2018. Đây cũng là bước tự kiểm tra, và cũng chính nó lôi ra "
        "được vấn đề tỷ lệ bê tông 2018 cao bất thường. Bài học nhóm rút ra là phần lớn giá trị của "
        "hai tuần vừa rồi nằm ở các bước tự nghi ngờ kết quả của chính mình, chứ không nằm ở bước chạy "
        "mô hình.")

    # --- SLIDE 7 ---
    SLIDE(doc, 7, "Chất lượng dữ liệu đầu vào", "1 phút")
    CUE(doc, "Slide: bảng số ảnh nguồn từng mốc — nhấn mạnh dòng 2018")
    TBL(doc, ["Mốc", "Khung mùa khô", "Số ảnh Sentinel-2"], [
        ["2018", "01/12/2017 – 30/04/2018", "3"],
        ["2020", "01/12/2019 – 30/04/2020", "52"],
        ["2022", "01/12/2021 – 30/04/2022", "45"],
        ["2024", "01/12/2023 – 30/04/2024", "51"],
        ["2026", "01/12/2025 – 30/04/2026", "46"],
    ])
    SAY(doc,
        "Trước khi chạy mô hình, nhóm đếm xem mỗi mốc có bao nhiêu ảnh nguồn để ghép. Kết quả có một "
        "điểm đáng chú ý: mốc 2018 chỉ có ba ảnh, trong khi các mốc khác đều từ 45 đến 52 ảnh.")
    SAY(doc,
        "Lý do là Sentinel-2 mới bắt đầu có dữ liệu Level-2A ổn định cho khu vực này từ cuối năm 2018. "
        "Ghép trung vị từ ba ảnh thì rõ ràng kém ổn định hơn ghép từ 50 ảnh. Nhóm ghi nhận điều này "
        "ngay từ đầu, và về sau nó đúng là nguyên nhân của một vấn đề mà em sẽ trình bày ở phần kết quả.")

    # --- SLIDE 8 ---
    SLIDE(doc, 8, "Lấy mẫu huấn luyện và một lỗi nhãn phát hiện được", "2,5 phút")
    CUE(doc, "Slide: bảng phân bố 4 lớp + ảnh contact sheet lớp Water")
    TBL(doc, ["Mã", "Lớp", "Số điểm ảnh trong vùng", "Tỷ lệ diện tích"], [
        ["0", "Water (mặt nước)", "1.002", "0,10%"],
        ["1", "Vegetation (thảm thực vật)", "979.874", "93,95%"],
        ["2", "Bare Soil (đất trống)", "18.659", "1,79%"],
        ["3", "Built-up (bề mặt xây dựng)", "43.464", "4,17%"],
    ])
    SAY(doc,
        "Nhãn huấn luyện lấy từ WorldCover, gộp 11 lớp gốc của họ về bốn lớp của đề tài. Nhóm lấy 300 "
        "điểm mẫu cho mỗi lớp, cố định số ngẫu nhiên bằng seed 42 để lần chạy sau ra đúng kết quả cũ, "
        "rồi chia 70 phần trăm huấn luyện và 30 phần trăm kiểm tra.")
    SAY(doc,
        "Nhưng nhóm không tin nhãn ngay. Bước tiếp theo là rà bằng mắt: lấy ngẫu nhiên 20 điểm mỗi lớp, "
        "cắt ảnh thật quanh từng điểm rồi nhìn xem nhãn có đúng không.")
    CUE(doc, "Chiếu contact sheet lớp Water, chỉ vào các ô số 6 đến 11")
    IMG(doc, GEEFIG / "qc_visual_water.png", "20 điểm mẫu lớp Water. Các ô 6–11 rơi vào giao lộ và khu dân cư, không phải mặt nước.", 5.2)
    SAY(doc,
        "Và đây là chỗ phát hiện vấn đề. Lớp Water có 6 trên 20 điểm, tức khoảng 30 phần trăm, rơi thẳng "
        "vào giao lộ hoặc khu dân cư chứ không phải mặt nước. Nhóm đoán nguyên nhân là mặt đường nhựa và "
        "mái nhà tối màu có phổ hơi giống nước nên WorldCover gán nhầm, hoặc là khu đó vốn là ao, đất "
        "ngập nước ở thời điểm WorldCover ghi nhận năm 2021, nhưng đã bị san lấp thành công trình ở thời "
        "điểm ảnh của nhóm.")
    SAY(doc,
        "Cách xử lý: lấy dư số điểm rồi lọc lại bằng ngưỡng chỉ số nước, vì mặt nước thật gần như luôn "
        "cho NDWI dương. Nhóm lấy 950 điểm, lọc còn 640 điểm đạt, tức loại đi 32,6 phần trăm.")
    SAY(doc,
        "Điểm hay ở đây là con số 32,6 phần trăm lọc bằng công thức rất khớp với 30 phần trăm mà nhóm "
        "đếm bằng mắt. Hai cách kiểm tra hoàn toàn độc lập nhưng ra cùng một kết quả, nên nhóm tin bước "
        "lọc này là đúng chứ không phải lọc bừa.")

    # --- SLIDE 9 ---
    SLIDE(doc, 9, "Huấn luyện mô hình: so sánh hai phương án", "2 phút")
    CUE(doc, "Slide: bảng so sánh A vs B")
    TBL(doc, ["Chỉ số", "Phương án A (9 đặc trưng)", "Phương án B (5 đặc trưng)"], [
        ["Overall Accuracy", "85,27%", "81,02%"],
        ["Hệ số Kappa", "0,803", "0,746"],
        ["Producer acc. — Water", "100,0%", "100,0%"],
        ["Producer acc. — Vegetation", "80,4%", "70,7%"],
        ["Producer acc. — Bare Soil", "91,9%", "89,2%"],
        ["Producer acc. — Built-up", "68,5%", "64,0%"],
    ])
    SAY(doc,
        "Mô hình nhóm dùng là Random Forest với 200 cây. Nhóm thử hai phương án đặc trưng để so sánh.")
    SAY(doc,
        "Phương án A dùng đủ chín đặc trưng. Phương án B rút gọn còn năm, bỏ bớt ba kênh khả kiến và "
        "NDWI. Lý do thử phương án B là vì ở bước khảo sát dữ liệu ban đầu, nhóm thấy ba kênh khả kiến "
        "tương quan với nhau rất cao, hệ số trên 0,9, còn cặp NDVI và NDWI thì gần như đối xứng nghịch, "
        "tương quan âm 0,97. Nghĩa là có khả năng thông tin bị dư thừa, bỏ bớt đi mô hình vẫn chạy tốt "
        "mà lại nhẹ hơn.")
    SAY(doc,
        "Nhưng thực nghiệm bác bỏ giả thuyết đó. Phương án A thắng trên mọi chỉ số: độ chính xác tổng "
        "thể cao hơn 4,25 điểm phần trăm, Kappa cao hơn 0,057. Và quan trọng nhất, lớp Built-up — chính "
        "là đối tượng nghiên cứu của đề tài — cũng chính xác hơn ở phương án A. Nên nhóm chọn A.")

    # --- SLIDE 10 ---
    SLIDE(doc, 10, "Ma trận nhầm lẫn: mô hình sai ở đâu", "1,5 phút")
    CUE(doc, "Slide: ma trận nhầm lẫn phương án A")
    TBL(doc, ["Thực tế \\ Dự đoán", "Water", "Vegetation", "Bare Soil", "Built-up"], [
        ["Water", "98", "0", "0", "0"],
        ["Vegetation", "0", "74", "4", "14"],
        ["Bare Soil", "0", "0", "68", "6"],
        ["Built-up", "0", "16", "12", "61"],
    ])
    SAY(doc,
        "Ma trận nhầm lẫn cho biết mô hình sai ở chỗ nào, chứ không chỉ sai bao nhiêu.")
    SAY(doc,
        "Lớp Water tách biệt tuyệt đối, 98 trên 98 điểm đúng hết. Nhưng chỗ này nhóm phải nói thẳng: "
        "con số 100 phần trăm đó cần diễn giải thận trọng, vì mẫu lớp Water đã được nhóm lọc trước bằng "
        "ngưỡng NDWI, mà NDWI lại chính là một đặc trưng đầu vào của mô hình. Nên nó phản ánh sự nhất "
        "quán của quy trình lọc nhiều hơn là năng lực phân biệt độc lập của mô hình.")
    SAY(doc,
        "Nhầm lẫn tập trung ở ba lớp còn lại: 16 điểm Built-up bị đoán thành Vegetation, 12 điểm "
        "Built-up bị đoán thành Bare Soil. Điều này khớp đúng với những gì nhóm thấy khi rà mẫu bằng "
        "mắt: ở một công trường đang thi công, ranh giới giữa đất đã san lấp và bề mặt đã đổ bê tông "
        "rất mờ, ngay cả người nhìn cũng khó phân định.")

    # --- SLIDE 11 ---
    SLIDE(doc, 11, "Kết quả: bản đồ phân loại 5 mốc", "2 phút")
    CUE(doc, "Chiếu hai bản đồ 2018 và 2024 cạnh nhau — đây là slide gây ấn tượng nhất, để lâu một chút")
    IMG(doc, MAPS / "classified_final_2018.png", "Bản đồ phân loại năm 2018.", 4.6)
    IMG(doc, MAPS / "classified_final_2024.png", "Bản đồ phân loại năm 2024. Khối đỏ ở giữa trùng hình dạng đường băng và sân đỗ.", 4.6)
    SAY(doc,
        "Đây là kết quả trực quan nhất của cả đề tài. Bên trái là năm 2018, bên phải là năm 2024. Màu "
        "xanh lá là thảm thực vật, màu đỏ là bề mặt xây dựng, màu be là đất trống, màu xanh dương là "
        "mặt nước.")
    SAY(doc,
        "Ở năm 2018, vùng lõi trung tâm còn là một mảng xanh liền, chính là đất nông nghiệp. Đến năm "
        "2024, đúng vùng đó biến thành một khối đỏ lớn, và hình dạng của khối đỏ này trùng khớp với "
        "đường băng và sân đỗ của sân bay.")
    SAY(doc,
        "Nhóm muốn nhấn mạnh: hình dạng đường băng này không phải nhóm vẽ vào, cũng không phải nhóm chỉ "
        "cho máy biết chỗ nào là sân bay. Mô hình chỉ được học từ đặc trưng phổ của 1.200 điểm mẫu, rồi "
        "tự nó phân loại ra hình dạng đó. Đây là bằng chứng mạnh cho thấy mô hình thực sự học được đúng "
        "thứ cần học.")

    # --- SLIDE 12 ---
    SLIDE(doc, 12, "Kết quả: số liệu diện tích qua 5 mốc", "2 phút")
    CUE(doc, "Slide: bảng diện tích + biểu đồ cột tỷ lệ Built-up qua các năm")
    TBL(doc, ["Mốc", "Built-up", "Vegetation", "Bare Soil", "Water"], [
        ["2018", "1.342,0 ha (12,93%)", "8.888,7 ha (85,61%)", "54,1 ha (0,52%)", "98,2 ha (0,95%)"],
        ["2020", "1.178,3 ha (11,30%)", "9.101,6 ha (87,28%)", "127,5 ha (1,22%)", "21,1 ha (0,20%)"],
        ["2022", "1.026,7 ha (9,84%)", "9.027,9 ha (86,57%)", "349,5 ha (3,35%)", "24,5 ha (0,23%)"],
        ["2024", "2.756,8 ha (26,44%)", "6.119,7 ha (58,68%)", "1.481,4 ha (14,20%)", "70,6 ha (0,68%)"],
        ["2026", "2.625,1 ha (25,17%)", "6.354,4 ha (60,93%)", "1.394,7 ha (13,37%)", "54,3 ha (0,52%)"],
    ])
    SAY(doc,
        "Còn đây là số liệu. Cột quan trọng nhất là cột Built-up.")
    SAY(doc,
        "Giai đoạn 2018 đến 2022, tỷ lệ bê tông hóa dao động quanh mức 10 đến 13 phần trăm, gần như "
        "không đổi. Nhưng sang 2024 thì nhảy vọt lên 26,4 phần trăm, tức tăng hơn gấp hai lần rưỡi chỉ "
        "trong hai năm, từ khoảng 1.000 héc ta lên gần 2.800 héc ta.")
    SAY(doc,
        "Con số này khớp chính xác với tiến độ thực tế của dự án. Sân bay khởi công đầu năm 2021, và "
        "phần lớn khối lượng xây dựng đường băng, nhà ga, đường công vụ diễn ra trong giai đoạn 2022 "
        "đến 2024. Đây là câu trả lời cho câu hỏi nghiên cứu thứ hai: tốc độ bê tông hóa nhanh nhất "
        "rơi vào giai đoạn 2022 đến 2024.")
    SAY(doc,
        "Một điểm nữa đáng chú ý là cột Bare Soil, đất trống. Nó tăng từ 0,5 phần trăm năm 2018 lên "
        "14,2 phần trăm năm 2024. Đất trống ở đây chính là mặt bằng đã giải phóng, đã san lấp nhưng "
        "chưa đổ bê tông. Nói cách khác, đây là bê tông của ngày mai.")

    # --- SLIDE 13 ---
    SLIDE(doc, 13, "Một vấn đề nhóm gặp phải và cách xử lý", "2 phút")
    CUE(doc, "Slide: bảng 3 dòng diễn biến xử lý 2018 + 2 ảnh bản đồ 2018 trước/sau")
    TBL(doc, ["Bước xử lý", "2018", "2020", "2022"], [
        ["Trước xử lý", "16,2%", "13,4%", "12,0%"],
        ["Sau lọc majority 3×3", "14,9%", "12,2%", "10,8%"],
        ["Sau khi bổ sung mẫu tay cho 2018", "12,9%", "11,3%", "9,8%"],
    ])
    SAY(doc,
        "Slide này nhóm muốn nói về một lỗi mà nhóm gặp, vì cách xử lý nó cũng là một phần kết quả.")
    SAY(doc,
        "Kết quả chạy lần đầu cho tỷ lệ bê tông hóa năm 2018 là 16,2 phần trăm, cao hơn cả 2020 và "
        "2022. Điều này vô lý, vì 2018 là mốc trước khi khởi công, lẽ ra phải thấp nhất.")
    CUE(doc, "Chiếu bản đồ 2018 trước xử lý")
    IMG(doc, MAPS / "classified_2018.png", "Bản đồ 2018 trước xử lý: nhiễu đốm đỏ rải khắp vùng thảm thực vật.", 4.4)
    SAY(doc,
        "Nhóm đối chiếu bản đồ phân loại với ảnh thật thì thấy: vùng lõi sân bay KHÔNG bị gán nhầm hàng "
        "loạt thành bê tông, đất nông nghiệp vẫn ra màu xanh đúng. Nhưng có nhiễu đốm đỏ rải rác khắp "
        "vùng thực vật, dạng muối tiêu. Nguyên nhân quay lại đúng chỗ nhóm đã ghi nhận ở slide 7: ảnh "
        "ghép 2018 chỉ dựng từ ba ảnh nguồn nên kém ổn định về phổ.")
    SAY(doc,
        "Nhóm xử lý hai bước. Bước một là lọc majority ba nhân ba, mỗi điểm ảnh được gán lại theo nhãn "
        "phổ biến nhất trong chín điểm lân cận. Đây là kỹ thuật hậu xử lý chuẩn trong viễn thám. Nó kéo "
        "2018 xuống 14,9 phần trăm nhưng chưa giải quyết được thứ tự bất hợp lý.")
    SAY(doc,
        "Bước hai là bổ sung mẫu. Nhóm chọn một vùng đã xác nhận bằng mắt chắc chắn là đất nông nghiệp, "
        "lấy 80 điểm trong đó, nhưng trích đặc trưng trực tiếp từ chính ảnh ghép 2018 chứ không phải từ "
        "ảnh tham chiếu 2022, gán nhãn thảm thực vật rồi huấn luyện lại. Sau bước này, 2018 xuống còn "
        "12,9 phần trăm, và độ chính xác trên tập kiểm tra không những không giảm mà còn nhích lên "
        "85,8 phần trăm với Kappa 0,811.")
    SAY(doc,
        "Nhóm xin nói thẳng một điều: sau xử lý thì 2018 vẫn còn cao hơn 2020 khoảng 1,6 điểm phần trăm. "
        "Nhóm không chỉnh thêm số liệu để ép nó thành một đường tăng đều đẹp mắt. Chênh lệch còn lại có "
        "thể là thật, do một số công trình nông thôn nhỏ bị dỡ khi giải phóng mặt bằng, mà cũng có thể "
        "vẫn còn dư sai lệch từ ảnh ghép 2018. Nhóm ghi rõ đây là hạn chế.")

    # --- SLIDE 14 ---
    SLIDE(doc, 14, "Kiểm chứng chéo bằng vệ tinh thứ hai", "1,5 phút")
    CUE(doc, "Slide: bảng hệ số tương quan Pearson 4 năm")
    TBL(doc, ["Mốc", "Hệ số tương quan r", "Phương sai giải thích (r²)", "Mức độ"], [
        ["2018", "0,714", "0,51", "Trung bình khá"],
        ["2020", "0,836", "0,70", "Khá"],
        ["2022", "0,733", "0,54", "Trung bình khá"],
        ["2024", "0,901", "0,81", "Cao"],
    ])
    SAY(doc,
        "Câu hỏi nghiên cứu thứ tư là kết quả có đáng tin không, hay chỉ là sản phẩm của một nguồn dữ "
        "liệu duy nhất. Để trả lời, nhóm tính chỉ số bê tông NDBI song song trên Landsat 8, một vệ tinh "
        "hoàn toàn khác, rồi đối chiếu với Sentinel-2.")
    SAY(doc,
        "Để so sánh được, nhóm phải hạ Sentinel-2 từ lưới 10 mét xuống lưới 30 mét cho khớp với Landsat, "
        "và bắt buộc dùng đúng cùng khung mùa khô. Nếu lấy khung cả năm thì hệ số tương quan sẽ phản ánh "
        "cả sai khác mùa vụ chứ không riêng sai khác giữa hai cảm biến.")
    SAY(doc,
        "Kết quả: tương quan từ 0,714 đến 0,901, trung bình trên 0,79. Hai vệ tinh nhất quán ở mức chấp "
        "nhận được. Nhóm cũng lưu ý rằng phải bình phương hệ số này mới ra phần phương sai giải thích "
        "được, nên 0,714 chỉ tương ứng khoảng 51 phần trăm, chứ không phải 71 phần trăm.")
    SAY(doc,
        "Và có một chi tiết rất đáng chú ý: mốc 2018 có hệ số tương quan thấp nhất. Điều này trùng khớp "
        "với phát hiện độc lập ở slide trước rằng ảnh ghép 2018 kém ổn định do chỉ có ba ảnh nguồn. Hai "
        "bằng chứng độc lập cùng chỉ về một mốc thời gian, nên nó củng cố nhận định của nhóm chứ không "
        "làm suy yếu kết quả chung.")

    # --- SLIDE 15 ---
    SLIDE(doc, 15, "Phân tích theo vành đai khoảng cách", "1,5 phút")
    CUE(doc, "Slide: sơ đồ 3 vành đai 0-2, 2-4, 4-6 km quanh tâm dự án")
    SAY(doc,
        "Câu hỏi nghiên cứu thứ ba là bê tông hóa có phụ thuộc khoảng cách tới sân bay không. Nhóm chia "
        "vùng nghiên cứu thành các vành đai theo khoảng cách tới tâm dự án rồi tính tỷ lệ bê tông trong "
        "từng vành.")
    SAY(doc,
        "Ở đây nhóm phải điều chỉnh kế hoạch ban đầu. Dự kiến ban đầu dùng vòng đệm 5, 10 và 15 ki lô "
        "mét. Nhưng khi kiểm tra lại mới thấy vùng nghiên cứu chỉ rộng 102 ki lô mét vuông, quá nhỏ để "
        "chứa các vòng đệm đó, và tâm dự án còn lệch về phía bắc nên cạnh bắc chỉ cách tâm 3,33 ki lô "
        "mét. Nhóm chuyển sang vành đai 0 đến 2, 2 đến 4 và 4 đến 6 ki lô mét.")
    SAY(doc,
        "Thêm nữa, các vành từ 4 ki lô mét trở ra không còn là vòng tròn đầy đủ mà chỉ là cung tròn "
        "thiếu phần phía bắc. Vì vậy nhóm so sánh bằng tỷ lệ phần trăm thay cho diện tích tuyệt đối, và "
        "trong code có hàm tính luôn phần vành đai bị cắt để nêu minh bạch trong báo cáo, thay vì chỉ "
        "nói chung chung là vành ngoài không đầy đủ.")

    # --- SLIDE 16 ---
    SLIDE(doc, 16, "Hạn chế của nghiên cứu", "1,5 phút")
    CUE(doc, "Slide: 4 gạch đầu dòng hạn chế — nói thẳng, đừng né")
    SAY(doc,
        "Nhóm xin nêu rõ bốn hạn chế, vì em nghĩ nêu ra thì đáng tin hơn là giấu đi.")
    SAY(doc,
        "Thứ nhất, độ chính xác 85,3 phần trăm được tính trên tập kiểm tra cân bằng, mỗi lớp 300 điểm. "
        "Nhưng ngoài thực tế thảm thực vật chiếm gần 94 phần trăm diện tích còn mặt nước chỉ 0,1 phần "
        "trăm. Nên con số 85,3 phần trăm không đại diện đúng cho độ chính xác thật của bản đồ. Nhóm đã "
        "bổ sung phép tính độ chính xác có trọng số theo diện tích theo khuyến nghị của Olofsson và "
        "cộng sự năm 2014, và đó mới là con số nên dùng trong báo cáo cuối.")
    SAY(doc,
        "Thứ hai, độ chính xác 100 phần trăm của lớp Water là do quy trình lọc chứ không phải năng lực "
        "phân biệt độc lập, như em đã nói ở slide ma trận nhầm lẫn.")
    SAY(doc,
        "Thứ ba, lớp Built-up là đối tượng chính của đề tài nhưng lại có độ chính xác thấp nhất, 68,5 "
        "phần trăm, do nhầm với đất trống và thực vật ở các khu đang thi công.")
    SAY(doc,
        "Thứ tư, tỷ lệ bê tông năm 2026 thấp hơn 2024 khoảng 1,3 điểm phần trăm. Chênh lệch này nằm "
        "trong biên độ sai số của mô hình, nhóm chưa đủ cơ sở để kết luận có suy giảm thật.")

    # --- SLIDE 17 ---
    SLIDE(doc, 17, "Kết luận", "1,5 phút")
    CUE(doc, "Slide: 4 kết luận đánh số, khớp với 4 câu hỏi nghiên cứu ở slide 3")
    SAY(doc,
        "Quay lại bốn câu hỏi đặt ra ở đầu bài, nhóm xin trả lời lần lượt.")
    SAY(doc,
        "Một. Đến năm 2026, khoảng 25,2 phần trăm diện tích vùng nghiên cứu đã bị bê tông hóa, tương "
        "đương 2.625 héc ta. So với mốc 2018 là 12,9 phần trăm thì diện tích đã tăng gần gấp đôi.")
    SAY(doc,
        "Hai. Tốc độ bê tông hóa nhanh nhất rơi vào giai đoạn 2022 đến 2024, tăng từ 9,8 lên 26,4 phần "
        "trăm, tức hơn gấp hai lần rưỡi chỉ trong hai năm. Giai đoạn 2018 đến 2022 gần như không đổi. "
        "Diễn biến này khớp với tiến độ thực tế của dự án.")
    SAY(doc,
        "Ba. Bê tông hóa tập trung rõ rệt ở vùng lõi dự án chứ không lan đều, thể hiện qua hình dạng "
        "khối đỏ trùng khớp đường băng trên bản đồ 2024, và qua phân tích vành đai khoảng cách.")
    SAY(doc,
        "Bốn. Về độ tin cậy, mô hình đạt độ chính xác tổng thể 85,8 phần trăm với Kappa 0,811, và kết "
        "quả được kiểm chứng chéo độc lập bằng vệ tinh Landsat 8 với tương quan trung bình trên 0,79. "
        "Nhóm cũng đã nêu rõ những chỗ còn hạn chế thay vì làm đẹp số liệu.")
    SAY(doc,
        "Về hướng phát triển, nếu có thêm thời gian nhóm sẽ làm ba việc: một là phân tích biến động "
        "từng điểm ảnh để biết chính xác đất chuyển từ loại gì sang loại gì, hai là bổ sung mẫu huấn "
        "luyện riêng cho từng mốc thời gian thay vì dùng chung một mốc tham chiếu, ba là mở rộng vùng "
        "nghiên cứu để phân tích vành đai được đầy đủ hơn.")

    # --- SLIDE 18 ---
    SLIDE(doc, 18, "Kết thúc", "15 giây")
    SAY(doc,
        "Phần trình bày của nhóm 07 đến đây là hết. Nhóm xin cảm ơn thầy và các bạn đã lắng nghe, và "
        "rất mong nhận được góp ý.")

    doc.add_page_break()

    # ================= PHẦN B =================
    H(doc, "PHẦN B — BẢNG SỐ LIỆU TRA NHANH", 1)
    P(doc, "Để mở sẵn khi bị hỏi con số cụ thể mà không nhớ.")

    H(doc, "B1. Thông số vùng nghiên cứu", 2)
    TBL(doc, ["Thông số", "Giá trị"], [
        ["Kích thước", "11,87 × 8,64 km (~102 km²)"],
        ["Giới hạn kinh độ", "106,9885° – 107,0970° Đông"],
        ["Giới hạn vĩ độ", "10,7368° – 10,8143° Bắc"],
        ["Tâm dự án", "107,04111° Đông; 10,78444° Bắc"],
        ["Ma trận ảnh @10 m", "1.190 cột × 870 hàng"],
        ["Dung lượng BSQ (19 kênh, 16-bit)", "1190 × 870 × 19 × 2 = 39.341.400 byte ≈ 39,34 MB"],
        ["Khoảng cách tâm tới cạnh bắc", "3,33 km"],
    ])

    H(doc, "B2. Cấu hình mô hình", 2)
    TBL(doc, ["Thông số", "Giá trị"], [
        ["Thuật toán", "Random Forest, 200 cây (ee.Classifier.smileRandomForest)"],
        ["Đặc trưng phương án A", "B2, B3, B4, B8, B11, B12, NDVI, NDBI, NDWI"],
        ["Đặc trưng phương án B", "B4, B8, B11, NDVI, NDBI"],
        ["Tổng số mẫu", "1.200 điểm (300/lớp)"],
        ["Chia train/test", "70/30 → 847 / 353 điểm"],
        ["Seed ngẫu nhiên", "42 (cố định để tái lập)"],
        ["Ảnh tham chiếu lấy mẫu", "Composite mùa khô 2022"],
        ["Hậu xử lý", "Lọc majority 3×3 (focalMode)"],
    ])

    H(doc, "B3. Công thức ba chỉ số phổ", 2)
    TBL(doc, ["Chỉ số", "Công thức", "Ý nghĩa"], [
        ["NDVI", "(B8 − B4) / (B8 + B4)", "Mật độ thảm thực vật; cây càng dày càng cao"],
        ["NDBI", "(B11 − B8) / (B11 + B8)", "Bề mặt xây dựng; bê tông, mái nhà cho giá trị cao"],
        ["NDWI", "(B3 − B8) / (B3 + B8)", "Mặt nước; nước gần như luôn cho giá trị dương"],
    ])

    H(doc, "B4. Bảng mã lớp và bảng màu", 2)
    TBL(doc, ["Mã", "Lớp", "Mã màu HEX"], [
        ["0", "Water — mặt nước", "#1E90FF (xanh dương)"],
        ["1", "Vegetation — thảm thực vật", "#2E8B57 (xanh lá đậm)"],
        ["2", "Bare Soil — đất trống", "#D2B48C (nâu nhạt)"],
        ["3", "Built-up — bề mặt xây dựng", "#D7301F (đỏ)"],
    ])

    H(doc, "B5. Bảng các lỗi đã phát hiện và cách xử lý", 2)
    P(doc, "Phần này rất dễ bị hỏi, nên nhớ ít nhất 3 dòng đầu.")
    TBL(doc, ["Vấn đề", "Cách xử lý"], [
        ["Tọa độ tâm sân bay ban đầu lệch 4–5 km về phía đông",
         "Xác định lại 107,04111°E / 10,78444°N, kiểm chứng bằng ảnh độ phân giải cao"],
        ["Ma trận ảnh tính sai, nhỏ hơn thực tế 10 lần mỗi chiều",
         "Bỏ phép chia dư thừa cho độ phân giải; kiểm chứng chéo bằng diện tích ÷ scale²"],
        ["Mỗi thành viên dùng một vùng nghiên cứu khác nhau",
         "Cố định một đa giác duy nhất bằng tọa độ tuyệt đối, dùng chung toàn nhóm"],
        ["Dùng lẫn Sentinel-2 Level-1C (chưa hiệu chỉnh khí quyển)",
         "Thống nhất chỉ dùng COPERNICUS/S2_SR_HARMONIZED mức Level-2A"],
        ["Khung thời gian lọc ảnh không thống nhất",
         "Thống nhất khung mùa khô 01/12 – 30/04 cho cả Sentinel-2 và Landsat"],
        ["~30% mẫu lớp Water của WorldCover rơi vào đường và mái nhà",
         "Lấy dư 950 điểm, lọc bằng NDWI > 0, loại 32,6% (khớp tỷ lệ đếm bằng mắt)"],
        ["Diện tích Built-up 2018 cao bất thường (16,2%)",
         "Lọc majority 3×3 + bổ sung 80 mẫu Vegetation trích từ chính composite 2018, huấn luyện lại"],
        ["Vùng nghiên cứu quá nhỏ so với vòng đệm 5/10/15 km",
         "Chuyển sang vành đai 0–2, 2–4, 4–6 km; so sánh bằng % thay vì diện tích tuyệt đối"],
    ])

    doc.add_page_break()

    # ================= PHẦN C =================
    H(doc, "PHẦN C — DỰ KIẾN CÂU HỎI PHẢN BIỆN VÀ CÁCH TRẢ LỜI", 1)

    qa = [
        ("Tại sao chọn Random Forest mà không phải mạng nơ-ron hay SVM?",
         "Ba lý do. Một, Random Forest chạy tốt với số mẫu vừa phải, 1.200 điểm là quá ít cho mạng "
         "nơ-ron nhưng đủ cho Random Forest. Hai, nó cho biết đặc trưng nào quan trọng, giúp nhóm giải "
         "thích được kết quả chứ không phải hộp đen. Ba, nó chạy sẵn trên Google Earth Engine nên không "
         "phải tải hàng chục GB ảnh về máy."),

        ("Vì sao lấy trung vị mà không lấy trung bình khi ghép ảnh?",
         "Vì trung bình bị kéo lệch bởi giá trị bất thường. Một điểm ảnh mây sót lại có giá trị phản xạ "
         "rất cao sẽ kéo trung bình lên, trong khi trung vị thì không bị ảnh hưởng nếu số điểm mây ít "
         "hơn một nửa. Đây là lý do trung vị là lựa chọn tiêu chuẩn khi ghép ảnh viễn thám."),

        ("Nhãn lấy từ WorldCover năm 2021, sao lại dùng để phân loại ảnh năm 2018 và 2026?",
         "Đây đúng là hạn chế lớn nhất về mặt thiết kế, nhóm thừa nhận. Mô hình học đặc trưng phổ của "
         "từng loại lớp phủ chứ không học vị trí, nên về nguyên tắc áp được sang mốc khác. Nhưng khi "
         "phổ của ảnh khác nhau nhiều thì mô hình chuyển kém, và mốc 2018 chính là ví dụ. Nhóm đã xử lý "
         "riêng cho 2018 bằng cách bổ sung mẫu trích trực tiếp từ ảnh 2018. Hướng làm đúng hơn là bổ "
         "sung mẫu riêng cho từng mốc, nhóm ghi vào phần hướng phát triển."),

        ("Tại sao độ chính xác lớp Water đạt 100% mà lại nói là không đáng tin?",
         "Vì nhóm đã lọc mẫu lớp này bằng ngưỡng NDWI dương trước khi huấn luyện, mà NDWI lại là một "
         "trong chín đặc trưng đầu vào của mô hình. Nói cách khác nhóm đã đưa cho mô hình một tập mẫu "
         "được chọn sẵn theo đúng tiêu chí mà mô hình dùng để phân biệt. Kết quả 100% phản ánh sự nhất "
         "quán của quy trình chứ không phải năng lực phân biệt độc lập."),

        ("Vì sao năm 2026 lại thấp hơn 2024? Bê tông không thể bị mất đi.",
         "Chênh lệch là 1,3 điểm phần trăm, nằm trong biên độ sai số của mô hình khi lớp Built-up chỉ "
         "đạt producer accuracy 68,5%. Nhóm không kết luận có suy giảm thật. Một khả năng nữa là sau "
         "khi công trình hoàn thiện, một phần diện tích được phủ cây xanh cảnh quan nên bị phân loại "
         "sang lớp thực vật, nhưng nhóm chưa kiểm chứng được điều này nên chỉ nêu là giả thuyết."),

        ("Vì sao không dùng ảnh độ phân giải cao hơn cho chính xác hơn?",
         "Ảnh độ phân giải dưới 1 mét đều là ảnh thương mại, phải trả phí, và không có chuỗi thời gian "
         "liên tục từ 2018 đến 2026 cho khu vực này. Sentinel-2 miễn phí, chụp lại mỗi 5 ngày, và ở độ "
         "phân giải 10 mét vẫn đủ để phân biệt bốn lớp lớp phủ ở quy mô công trình lớn như sân bay."),

        ("Kết quả này khác gì so với việc chỉ nhìn ảnh vệ tinh bằng mắt?",
         "Nhìn bằng mắt thì chỉ nói được là 'có vẻ xây nhiều hơn'. Còn cách làm này cho ra con số cụ "
         "thể tính bằng héc ta, lặp lại được, và kiểm chứng được độ chính xác bằng ma trận nhầm lẫn. "
         "Quan trọng hơn là quy trình này chạy lại được cho khu vực khác hoặc mốc thời gian khác mà "
         "không cần làm lại từ đầu."),

        ("Vì sao mốc 2018 chỉ có 3 ảnh mà vẫn giữ trong nghiên cứu?",
         "Vì 2018 là mốc nền trước khi khởi công, bỏ đi thì mất điểm so sánh gốc. Nhóm chọn cách giữ "
         "lại nhưng nêu rõ hạn chế và xử lý riêng, đồng thời có hai bằng chứng độc lập cùng chỉ ra "
         "mốc này kém ổn định là nhiễu đốm trên bản đồ và hệ số tương quan Landsat thấp nhất."),

        ("Nhóm có kiểm chứng thực địa không?",
         "Không, nhóm không đi thực địa được. Thay vào đó nhóm dùng ba cách kiểm chứng gián tiếp: rà "
         "mẫu bằng mắt trên ảnh, đối chiếu với vệ tinh thứ hai là Landsat 8, và đối chiếu xu hướng với "
         "mốc thời gian thực tế của dự án theo Quyết định 1777. Đây là hạn chế nhóm ghi nhận."),
    ]
    for i, (q, a) in enumerate(qa, 1):
        P(doc, f"Câu {i}. {q}", bold=True, space_after=3)
        P(doc, "Trả lời: " + a, space_after=10)

    doc.save(OUT)
    print(f"[done] Da ghi {OUT}")


if __name__ == "__main__":
    main()
