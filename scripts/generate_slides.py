#!/usr/bin/env python3
"""
Sinh bo slide .pptx (20 slide, 16:9) cho bai thuyet trinh de tai
"Do luong toc do be tong hoa khu vuc san bay Long Thanh 2018-2026".

Anh that duoc chen truc tiep tu outputs/. Moi con so lay tu ket qua da chay.

Chay:
    python scripts/generate_slides.py
"""

from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = PROJECT_ROOT / "reports" / "Slide_Thuyet_trinh_LongThanh.pptx"
FIG = PROJECT_ROOT / "outputs" / "figures"
GEEFIG = FIG / "gee"
MAPS = PROJECT_ROOT / "outputs" / "maps"

# Bang mau khop voi ban do trong bai
BLUE = RGBColor(0x1E, 0x90, 0xFF)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
TAN = RGBColor(0xD2, 0xB4, 0x8C)
RED = RGBColor(0xD7, 0x30, 0x1F)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Arial"
W, H = Inches(13.333), Inches(7.5)


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def textbox(slide, left, top, width, height, text, size=20, bold=False,
            color=BLACK, align=PP_ALIGN.LEFT, italic=False, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = FONT
    return tb


def header(slide, title, accent=RED):
    """Thanh tieu de + gach mau duoi."""
    textbox(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9),
            title, size=32, bold=True, color=BLACK)
    bar = slide.shapes.add_shape(1, Inches(0.6), Inches(1.28), Inches(1.6), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    bar.shadow.inherit = False


def bullets(slide, items, top=Inches(1.75), left=Inches(0.85), width=Inches(11.6), size=21, gap=0.62):
    for i, item in enumerate(items):
        y = top + Inches(i * gap)
        dot = slide.shapes.add_shape(9, left, y + Inches(0.11), Inches(0.13), Inches(0.13))
        dot.fill.solid()
        dot.fill.fore_color.rgb = RED
        dot.line.fill.background()
        dot.shadow.inherit = False
        textbox(slide, left + Inches(0.32), y, width, Inches(0.55), item, size=size)


def table(slide, headers, rows, left=Inches(0.7), top=Inches(1.7),
          width=Inches(11.9), height=None, fsize=14, hsize=14):
    nrows, ncols = len(rows) + 1, len(headers)
    height = height or Inches(min(0.42 * nrows, 5.2))
    shape = slide.shapes.add_table(nrows, ncols, left, top, width, height)
    tbl = shape.table
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = str(h)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(hsize)
                r.font.bold = True
                r.font.name = FONT
                r.font.color.rgb = WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x33)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(fsize)
                    r.font.name = FONT
                    r.font.color.rgb = BLACK
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else RGBColor(0xF2, 0xF2, 0xF2)
    return tbl


def picture(slide, path: Path, left, top, width=None, height=None):
    if not path.exists():
        textbox(slide, left, top, Inches(5), Inches(0.5),
                f"[Thiếu hình: {path.name}]", size=14, italic=True, color=GRAY)
        return None
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def caption(slide, text, left, top, width, size=12):
    textbox(slide, left, top, width, Inches(0.4), text, size=size,
            italic=True, color=GRAY, align=PP_ALIGN.CENTER)


def legend(slide, top=Inches(6.75), left=Inches(3.6)):
    items = [("Nước", BLUE), ("Thảm thực vật", GREEN), ("Đất trống", TAN), ("Bề mặt xây dựng", RED)]
    x = left
    for name, color in items:
        sw = slide.shapes.add_shape(1, x, top + Inches(0.05), Inches(0.22), Inches(0.22))
        sw.fill.solid()
        sw.fill.fore_color.rgb = color
        sw.line.fill.background()
        sw.shadow.inherit = False
        textbox(slide, x + Inches(0.3), top, Inches(1.9), Inches(0.35), name, size=13)
        x += Inches(2.1)


def note(slide, text, top=Inches(6.55), size=15, color=RED):
    textbox(slide, Inches(0.85), top, Inches(11.6), Inches(0.6), text, size=size, bold=True, color=color)


def build():
    prs = new_deck()

    # ---------- 1. TITLE ----------
    s = blank(prs)
    band = s.shapes.add_shape(1, Inches(0), Inches(0), W, Inches(0.22))
    band.fill.solid(); band.fill.fore_color.rgb = RED
    band.line.fill.background(); band.shadow.inherit = False
    textbox(s, Inches(1.0), Inches(2.15), Inches(11.3), Inches(1.9),
            "ĐO LƯỜNG TỐC ĐỘ BÊ TÔNG HÓA\nKHU VỰC SÂN BAY LONG THÀNH",
            size=42, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, Inches(1.0), Inches(4.05), Inches(11.3), Inches(0.6),
            "Giai đoạn 2018 – 2026", size=26, color=RED, align=PP_ALIGN.CENTER)
    textbox(s, Inches(1.0), Inches(5.1), Inches(11.3), Inches(1.2),
            "Nhóm 07  •  Học phần Xử lý ảnh\nThành viên: ………………………………………",
            size=17, color=GRAY, align=PP_ALIGN.CENTER)

    # ---------- 2. LÝ DO ----------
    s = blank(prs); header(s, "Lý do chọn đề tài")
    bullets(s, [
        "Công trường xây dựng lớn nhất Việt Nam hiện nay",
        "Phê duyệt đầu tư 2020 (QĐ 1777) → khởi công 2021 → khai trương 12/2025",
        "Toàn bộ quá trình nằm gọn trong khoảng thời gian Sentinel-2 có dữ liệu",
        "Bài toán đo được: bề mặt không thấm nước có đặc trưng phổ riêng biệt",
        "Rẻ hơn khảo sát thực địa, quy trình lặp lại được cho khu vực khác",
    ], top=Inches(2.0), gap=0.78, size=22)

    # ---------- 3. MỤC TIÊU ----------
    s = blank(prs); header(s, "Mục tiêu — 4 câu hỏi nghiên cứu")
    qs = [
        ("1", "Đến 2026, bao nhiêu % diện tích đã bị bê tông hóa?"),
        ("2", "Tốc độ bê tông hóa nhanh nhất rơi vào giai đoạn nào?"),
        ("3", "Mức độ bê tông hóa có phụ thuộc khoảng cách tới sân bay?"),
        ("4", "Kết quả có đáng tin cậy không?"),
    ]
    for i, (num, q) in enumerate(qs):
        y = Inches(2.0 + i * 1.15)
        box = s.shapes.add_shape(1, Inches(0.85), y, Inches(0.75), Inches(0.75))
        box.fill.solid(); box.fill.fore_color.rgb = RED
        box.line.fill.background(); box.shadow.inherit = False
        tf = box.text_frame; tf.text = num
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(26); r.font.bold = True
                r.font.color.rgb = WHITE; r.font.name = FONT
        textbox(s, Inches(1.85), y + Inches(0.12), Inches(10.8), Inches(0.7), q, size=22)

    # ---------- 4. VÙNG NGHIÊN CỨU ----------
    s = blank(prs); header(s, "Vùng nghiên cứu")
    picture(s, GEEFIG / "aoi_rgb_2022.png", Inches(0.7), Inches(1.65), height=Inches(4.6))
    caption(s, "Ảnh tổ hợp màu tự nhiên, composite mùa khô 2022", Inches(0.7), Inches(6.35), Inches(6.4))
    facts = [
        "Kích thước: 11,87 × 8,64 km (~102 km²)",
        "Kinh độ: 106,9885° – 107,0970° Đông",
        "Vĩ độ: 10,7368° – 10,8143° Bắc",
        "Tâm dự án: 107,04111°E / 10,78444°N",
        "Ma trận ảnh @10 m: 1.190 × 870 pixel",
    ]
    for i, f in enumerate(facts):
        textbox(s, Inches(7.6), Inches(1.85 + i * 0.55), Inches(5.3), Inches(0.5), "• " + f, size=18)
    textbox(s, Inches(7.6), Inches(4.8), Inches(5.3), Inches(1.6),
            "Cố định bằng tọa độ tuyệt đối, không khoanh tay\n"
            "→ 5 mốc khớp khít nhau tới từng điểm ảnh",
            size=17, bold=True, color=RED)

    # ---------- 5. DỮ LIỆU ----------
    s = blank(prs); header(s, "Dữ liệu sử dụng")
    table(s, ["Nguồn", "Mã bộ dữ liệu", "Độ phân giải", "Vai trò"], [
        ["Sentinel-2 Level-2A", "COPERNICUS/S2_SR_HARMONIZED", "10 m", "Nguồn chính"],
        ["Landsat 8 Collection 2 L2", "LANDSAT/LC08/C02/T1_L2", "30 m", "Kiểm chứng chéo"],
        ["ESA WorldCover v200", "ESA/WorldCover/v200", "10 m", "Nhãn huấn luyện (nền 2021)"],
    ], top=Inches(2.1), fsize=17, hsize=17, height=Inches(2.4))
    note(s, "Cả ba nguồn đều miễn phí và mở. Đã thống nhất chỉ dùng Level-2A "
            "(đã hiệu chỉnh khí quyển) cho toàn nhóm.", top=Inches(5.0), size=17, color=BLACK)

    # ---------- 6. QUY TRÌNH TỔNG THỂ ----------
    s = blank(prs); header(s, "Quy trình tổng thể")
    steps = ["Ảnh thô\nSentinel-2", "Ghép mùa khô\n+ che mây", "Bộ 9\nđặc trưng",
             "Lấy mẫu +\nlàm sạch nhãn", "Random Forest\n(A vs B)",
             "Phân loại 5 mốc\n+ lọc nhiễu", "Diện tích, vành đai,\nkiểm chứng"]
    colors = [GRAY, BLUE, GREEN, TAN, RED, RED, BLACK]
    x = Inches(0.45)
    for i, (st, col) in enumerate(zip(steps, colors)):
        box = s.shapes.add_shape(5, x, Inches(3.0), Inches(1.55), Inches(1.5))
        box.fill.solid(); box.fill.fore_color.rgb = col
        box.line.fill.background(); box.shadow.inherit = False
        tf = box.text_frame; tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for j, line in enumerate(st.split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = line
            r.font.size = Pt(13); r.font.bold = True
            r.font.color.rgb = WHITE; r.font.name = FONT
        if i < len(steps) - 1:
            textbox(s, x + Inches(1.55), Inches(3.5), Inches(0.35), Inches(0.5),
                    "›", size=26, bold=True, color=GRAY, align=PP_ALIGN.CENTER)
        x += Inches(1.83)
    note(s, "Toàn bộ 7 bước nằm trong một script GEE duy nhất, cố định seed = 42 → chạy lại ra đúng kết quả.",
         top=Inches(5.3), color=BLACK, size=17)

    # ---------- 7. PIPELINE 11 BƯỚC ----------
    s = blank(prs); header(s, "Pipeline chi tiết — 11 bước theo đúng code")
    table(s, ["#", "Nội dung"], [
        ["1", "Khai báo AOI + hằng số dùng chung (seed 42, ngưỡng mây 60%)"],
        ["2", "Ghép Sentinel-2 mùa khô: lọc cảnh → mask mây SCL (3/8/9/10) → ÷10.000 → median → clip"],
        ["3", "Tạo bộ 9 đặc trưng: 6 kênh phổ + NDVI + NDBI + NDWI"],
        ["4", "Lấy mẫu WorldCover: remap 11→4 lớp, 300 điểm/lớp, lọc Water bằng NDWI > 0"],
        ["5", "Huấn luyện Random Forest 200 cây — 2 phương án A (9 band) và B (5 band)"],
        ["6", "Độ chính xác có trọng số theo diện tích (Olofsson và cộng sự, 2014)"],
        ["7", "Áp mô hình thắng lên 5 mốc + lọc majority 3×3 khử nhiễu muối tiêu"],
        ["8", "Tính diện tích từng lớp bằng ee.Image.pixelArea()"],
        ["9", "Vành đai 0–2, 2–4, 4–6 km + định lượng phần vành bị cắt"],
        ["10", "Kiểm chứng Landsat 8: hiệu chỉnh C2, mask QA_PIXEL, hạ S2 về 30 m, Pearson"],
        ["11", "Xuất kết quả ra Google Drive"],
    ], top=Inches(1.6), fsize=13.5, hsize=14, height=Inches(5.4), width=Inches(11.9))

    # ---------- 8. CHIA VIỆC 3.1-3.9 ----------
    s = blank(prs); header(s, "Cách chia việc — 9 đầu việc tuần 2")
    table(s, ["Task", "Nội dung", "Kết quả thực tế"], [
        ["3.1", "Chuyển pipeline AWS/STAC → GEE", "Script GEE thống nhất"],
        ["3.2", "Tạo bộ đặc trưng 9 band", "Ảnh 9 kênh cho 5 mốc"],
        ["3.3", "Lấy mẫu WorldCover, 300 điểm/lớp, chia 70/30", "1.200 điểm — train 847 / test 353"],
        ["3.4", "Rà mẫu bằng mắt ~20 điểm/lớp", "Phát hiện lớp Water sai ~30%"],
        ["3.5", "Train Random Forest phương án A", "OA 85,27% — Kappa 0,803"],
        ["3.6", "Train phương án B (rút gọn)", "OA 81,02% — Kappa 0,746"],
        ["3.7", "So sánh A vs B", "Chọn A trên mọi chỉ số"],
        ["3.8", "Áp mô hình lên 5 mốc", "5 bản đồ + bảng diện tích"],
        ["3.9", "Kiểm chứng riêng mốc 2018", "Phát hiện & xử lý bất thường 16,2%"],
    ], top=Inches(1.6), fsize=14, hsize=14.5, height=Inches(4.8))
    note(s, "Giá trị lớn nhất nằm ở task 3.4 và 3.9 — hai bước tự nghi ngờ kết quả của chính mình.",
         top=Inches(6.6))

    # ---------- 9. CHẤT LƯỢNG DỮ LIỆU ----------
    s = blank(prs); header(s, "Chất lượng dữ liệu đầu vào")
    table(s, ["Mốc", "Khung mùa khô", "Số ảnh Sentinel-2"], [
        ["2018", "01/12/2017 – 30/04/2018", "3"],
        ["2020", "01/12/2019 – 30/04/2020", "52"],
        ["2022", "01/12/2021 – 30/04/2022", "45"],
        ["2024", "01/12/2023 – 30/04/2024", "51"],
        ["2026", "01/12/2025 – 30/04/2026", "46"],
    ], top=Inches(1.9), fsize=17, hsize=17, height=Inches(3.2), width=Inches(9.0), left=Inches(2.1))
    note(s, "Mốc 2018 chỉ có 3 ảnh — Sentinel-2 mới ổn định Level-2A cho khu vực này từ cuối 2018.\n"
            "Đây là nguyên nhân gốc của vấn đề được xử lý riêng ở phần sau.", top=Inches(5.5), size=17)

    # ---------- 10. PHÂN BỐ 4 LỚP ----------
    s = blank(prs); header(s, "Phân bố 4 lớp lớp phủ trong vùng nghiên cứu")
    table(s, ["Mã", "Lớp", "Số điểm ảnh", "Tỷ lệ diện tích"], [
        ["0", "Water — mặt nước", "1.002", "0,10%"],
        ["1", "Vegetation — thảm thực vật", "979.874", "93,95%"],
        ["2", "Bare Soil — đất trống", "18.659", "1,79%"],
        ["3", "Built-up — bề mặt xây dựng", "43.464", "4,17%"],
    ], top=Inches(2.1), fsize=18, hsize=18, height=Inches(2.8), width=Inches(10.0), left=Inches(1.6))
    note(s, "Mất cân bằng rất mạnh: Vegetation chiếm gần 94%, Water chỉ 0,10% (1.002 điểm ảnh).",
         top=Inches(5.5), size=18)

    # ---------- 11. LỖI NHÃN WATER ----------
    s = blank(prs); header(s, "Lỗi nhãn phát hiện được khi rà bằng mắt")
    picture(s, GEEFIG / "qc_visual_water.png", Inches(0.6), Inches(1.6), height=Inches(4.6))
    caption(s, "20 điểm mẫu lớp Water — ô 6–11 rơi vào giao lộ, khu dân cư",
            Inches(0.6), Inches(6.3), Inches(6.6))
    pts = [
        "Rà 20 điểm/lớp trên ảnh thật",
        "Lớp Water sai 6/20 → ~30%",
        "Nguyên nhân: đường nhựa và mái nhà tối màu\nphổ gần giống nước",
        "Xử lý: lấy dư 950 điểm, lọc NDWI > 0\n→ loại 32,6%",
    ]
    y = 1.9
    for p in pts:
        textbox(s, Inches(7.5), Inches(y), Inches(5.4), Inches(0.9), "• " + p, size=18)
        y += 1.0
    textbox(s, Inches(7.5), Inches(5.9), Inches(5.4), Inches(1.0),
            "32,6% (lọc bằng công thức)\n≈ 30% (đếm bằng mắt)\n→ hai cách độc lập, cùng kết quả",
            size=17, bold=True, color=RED)

    # ---------- 12. SO SÁNH A vs B ----------
    s = blank(prs); header(s, "So sánh hai phương án đặc trưng")
    table(s, ["Chỉ số", "Phương án A (9 đặc trưng)", "Phương án B (5 đặc trưng)"], [
        ["Overall Accuracy", "85,27%", "81,02%"],
        ["Hệ số Kappa", "0,803", "0,746"],
        ["Producer acc. — Water", "100,0%", "100,0%"],
        ["Producer acc. — Vegetation", "80,4%", "70,7%"],
        ["Producer acc. — Bare Soil", "91,9%", "89,2%"],
        ["Producer acc. — Built-up", "68,5%", "64,0%"],
    ], top=Inches(1.9), fsize=17, hsize=17, height=Inches(3.6), width=Inches(10.6), left=Inches(1.35))
    note(s, "→ Chọn phương án A: hơn 4,25 điểm % về OA, và lớp Built-up (đối tượng chính) cũng tốt hơn.",
         top=Inches(6.0), size=18)

    # ---------- 13. MA TRẬN NHẦM LẪN ----------
    s = blank(prs); header(s, "Ma trận nhầm lẫn — mô hình sai ở đâu")
    table(s, ["Thực tế \\ Dự đoán", "Water", "Vegetation", "Bare Soil", "Built-up"], [
        ["Water", "98", "0", "0", "0"],
        ["Vegetation", "0", "74", "4", "14"],
        ["Bare Soil", "0", "0", "68", "6"],
        ["Built-up", "0", "16", "12", "61"],
    ], top=Inches(2.0), fsize=18, hsize=17, height=Inches(2.6), width=Inches(10.4), left=Inches(1.45))
    bullets(s, [
        "Water tách biệt tuyệt đối — nhưng do đã lọc trước bằng NDWI, không phải năng lực độc lập",
        "Nhầm lẫn tập trung: Built-up ↔ Vegetation (16 điểm) và Built-up ↔ Bare Soil (12 điểm)",
        "Đúng với thực tế: ở công trường, ranh giới đất san lấp / đã đổ bê tông rất mờ",
    ], top=Inches(5.1), size=17, gap=0.62)

    # ---------- 14. BẢN ĐỒ 2018 vs 2024 ----------
    s = blank(prs); header(s, "Kết quả: bản đồ phân loại 2018 và 2024")
    picture(s, MAPS / "classified_final_2018.png", Inches(0.75), Inches(1.7), height=Inches(4.1))
    picture(s, MAPS / "classified_final_2024.png", Inches(6.95), Inches(1.7), height=Inches(4.1))
    caption(s, "Năm 2018 — vùng lõi còn là đất nông nghiệp", Inches(0.75), Inches(5.9), Inches(5.7), size=14)
    caption(s, "Năm 2024 — khối đỏ trùng hình dạng đường băng", Inches(6.95), Inches(5.9), Inches(5.7), size=14)
    legend(s, top=Inches(6.45), left=Inches(2.6))
    note(s, "Hình dạng đường băng do mô hình tự phân loại ra, không phải nhóm vẽ vào.",
         top=Inches(6.95), size=15)

    # ---------- 15. BẢNG DIỆN TÍCH ----------
    s = blank(prs); header(s, "Kết quả: diện tích 4 lớp qua 5 mốc")
    table(s, ["Mốc", "Built-up", "Vegetation", "Bare Soil", "Water"], [
        ["2018", "1.342,0 ha (12,93%)", "8.888,7 ha (85,61%)", "54,1 ha (0,52%)", "98,2 ha (0,95%)"],
        ["2020", "1.178,3 ha (11,30%)", "9.101,6 ha (87,28%)", "127,5 ha (1,22%)", "21,1 ha (0,20%)"],
        ["2022", "1.026,7 ha (9,84%)", "9.027,9 ha (86,57%)", "349,5 ha (3,35%)", "24,5 ha (0,23%)"],
        ["2024", "2.756,8 ha (26,44%)", "6.119,7 ha (58,68%)", "1.481,4 ha (14,20%)", "70,6 ha (0,68%)"],
        ["2026", "2.625,1 ha (25,17%)", "6.354,4 ha (60,93%)", "1.394,7 ha (13,37%)", "54,3 ha (0,52%)"],
    ], top=Inches(1.75), fsize=15, hsize=15.5, height=Inches(3.0))
    bullets(s, [
        "2018–2022: dao động quanh 10–13%, gần như không đổi",
        "2022 → 2024: nhảy vọt 9,84% → 26,44%, tăng hơn 2,5 lần chỉ trong 2 năm",
        "Bare Soil tăng 0,52% → 14,20%: mặt bằng đã san lấp, chính là “bê tông của ngày mai”",
    ], top=Inches(5.05), size=18, gap=0.66)

    # ---------- 16. XỬ LÝ 2018 ----------
    s = blank(prs); header(s, "Xử lý bất thường ở mốc 2018")
    picture(s, MAPS / "classified_2018.png", Inches(0.65), Inches(1.75), height=Inches(3.3))
    caption(s, "2018 trước xử lý — nhiễu đốm đỏ rải khắp vùng thực vật",
            Inches(0.65), Inches(5.15), Inches(5.3), size=13)
    table(s, ["Bước xử lý", "2018", "2020", "2022"], [
        ["Trước xử lý", "16,2%", "13,4%", "12,0%"],
        ["Sau lọc majority 3×3", "14,9%", "12,2%", "10,8%"],
        ["Sau bổ sung 80 mẫu tay cho 2018", "12,9%", "11,3%", "9,8%"],
    ], top=Inches(1.9), left=Inches(6.5), width=Inches(6.2), fsize=14, hsize=14, height=Inches(1.9))
    textbox(s, Inches(6.5), Inches(4.1), Inches(6.2), Inches(2.2),
            "• 16,2% cao hơn cả 2020 và 2022 → vô lý\n"
            "• Nguyên nhân: composite 2018 chỉ có 3 ảnh nguồn\n"
            "• Sau xử lý: OA tăng lên 85,8% — Kappa 0,811\n"
            "• Vẫn còn cao hơn 2020 ~1,6 điểm % → ghi nhận là hạn chế,\n"
            "   không chỉnh số cho đẹp",
            size=16)

    # ---------- 17. KIỂM CHỨNG LANDSAT ----------
    s = blank(prs); header(s, "Kiểm chứng chéo bằng vệ tinh Landsat 8")
    table(s, ["Mốc", "Hệ số tương quan r", "Phương sai giải thích (r²)", "Mức độ"], [
        ["2018", "0,714", "0,51", "Trung bình khá"],
        ["2020", "0,836", "0,70", "Khá"],
        ["2022", "0,733", "0,54", "Trung bình khá"],
        ["2024", "0,901", "0,81", "Cao"],
    ], top=Inches(1.95), fsize=18, hsize=17, height=Inches(2.7), width=Inches(10.2), left=Inches(1.55))
    bullets(s, [
        "Tính NDBI song song trên 2 vệ tinh, hạ Sentinel-2 từ 10 m về 30 m để khớp lưới",
        "Tương quan trung bình > 0,79 → hai nguồn nhất quán ở mức chấp nhận được",
        "Mốc 2018 thấp nhất — trùng khớp với phát hiện độc lập rằng ảnh ghép 2018 kém ổn định",
    ], top=Inches(5.1), size=17, gap=0.62)

    # ---------- 18. VÀNH ĐAI ----------
    s = blank(prs); header(s, "Phân tích theo vành đai khoảng cách")
    cx, cy = Inches(3.4), Inches(4.1)
    for r_in, col in [(2.45, TAN), (1.65, GREEN), (0.85, RED)]:
        c = s.shapes.add_shape(9, cx - Inches(r_in), cy - Inches(r_in),
                               Inches(2 * r_in), Inches(2 * r_in))
        c.fill.background()
        c.line.color.rgb = col
        c.line.width = Pt(2.5)
        c.shadow.inherit = False
    dot = s.shapes.add_shape(9, cx - Inches(0.08), cy - Inches(0.08), Inches(0.16), Inches(0.16))
    dot.fill.solid(); dot.fill.fore_color.rgb = BLACK
    dot.line.fill.background(); dot.shadow.inherit = False
    textbox(s, cx + Inches(0.15), cy - Inches(0.45), Inches(2.0), Inches(0.4), "Tâm dự án", size=13)
    for i, lbl in enumerate(["0–2 km", "2–4 km", "4–6 km"]):
        textbox(s, cx + Inches(0.55 + i * 0.8), cy + Inches(0.35 + i * 0.75),
                Inches(1.6), Inches(0.35), lbl, size=14, bold=True)
    bullets(s, [
        "Kế hoạch ban đầu: vòng đệm 5 / 10 / 15 km",
        "Thực tế AOI chỉ 102 km², cạnh bắc cách tâm 3,33 km",
        "→ Đổi sang vành đai 0–2, 2–4, 4–6 km",
        "Vành từ 4 km trở ra chỉ là cung tròn, thiếu phần phía bắc",
        "→ So sánh bằng % thay cho diện tích tuyệt đối",
        "Code tính luôn phần vành bị cắt để nêu minh bạch",
    ], top=Inches(1.85), left=Inches(7.0), width=Inches(5.9), size=17, gap=0.72)

    # ---------- 19. HẠN CHẾ ----------
    s = blank(prs); header(s, "Hạn chế của nghiên cứu")
    bullets(s, [
        "OA 85,3% tính trên tập test cân bằng lớp — chưa phản ánh đúng thực tế\n"
        "(Vegetation 93,95% vs Water 0,10%) → đã bổ sung OA có trọng số diện tích",
        "Water đạt 100% là do quy trình lọc NDWI, không phải năng lực phân biệt độc lập",
        "Built-up là đối tượng chính nhưng producer accuracy thấp nhất: 68,5%",
        "2026 thấp hơn 2024 1,3 điểm % — nằm trong biên độ sai số, chưa kết luận được",
    ], top=Inches(1.95), size=19, gap=1.15)

    # ---------- 20. KẾT LUẬN ----------
    s = blank(prs); header(s, "Kết luận")
    answers = [
        ("1", "25,17% (~2.625 ha) đã bê tông hóa vào 2026 — gần gấp đôi 2018 (12,93%)"),
        ("2", "Nhanh nhất 2022 → 2024: 9,84% → 26,44%, hơn 2,5 lần trong 2 năm"),
        ("3", "Tập trung ở vùng lõi dự án, không lan đều ra xung quanh"),
        ("4", "OA 85,8% — Kappa 0,811, kiểm chứng chéo Landsat r trung bình > 0,79"),
    ]
    for i, (num, a) in enumerate(answers):
        y = Inches(1.85 + i * 0.95)
        box = s.shapes.add_shape(1, Inches(0.85), y, Inches(0.62), Inches(0.62))
        box.fill.solid(); box.fill.fore_color.rgb = RED
        box.line.fill.background(); box.shadow.inherit = False
        tf = box.text_frame; tf.text = num
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(22); r.font.bold = True
                r.font.color.rgb = WHITE; r.font.name = FONT
        textbox(s, Inches(1.7), y + Inches(0.08), Inches(11.0), Inches(0.7), a, size=19)
    textbox(s, Inches(0.85), Inches(5.85), Inches(11.9), Inches(1.3),
            "Hướng phát triển: phân tích biến động từng điểm ảnh  •  "
            "bổ sung mẫu riêng cho từng mốc  •  mở rộng vùng nghiên cứu",
            size=17, bold=True, color=GRAY)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"[done] Da ghi {OUT} ({len(prs.slides.__iter__.__self__._sldIdLst)} slide)")


if __name__ == "__main__":
    build()
