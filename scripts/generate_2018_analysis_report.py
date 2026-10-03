#!/usr/bin/env python3
"""
Sinh bao cao phan tich rieng cho anh Sentinel-2 nam 2018 (baseline truoc khi
xay san bay Long Thanh), dung Times New Roman, chu mau den, chen hinh that
tu outputs/figures/*_2018.png.

Chay:
    python scripts/generate_2018_analysis_report.py
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

from src import config as cfg
from src.report_images import optimized
from src import validation, preprocessing, features

BLACK = RGBColor(0, 0, 0)
FONT_NAME = "Times New Roman"
OUT_PATH = cfg.REPORTS_DIR / "Phan_tich_du_lieu_2018.docx"


def set_run_font(run, size=13, bold=False, italic=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic
    # ép luôn font cho phần chữ Đông Á / phức hợp, tránh Word tự đổi font khi gặp dấu
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT_NAME)


def add_heading(doc, text, level=1):
    h = doc.add_heading("", level=level)
    run = h.add_run(text)
    size = {1: 16, 2: 14, 3: 13}.get(level, 13)
    set_run_font(run, size=size, bold=True)
    return h


def add_para(doc, text, size=13, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=12, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            set_run_font(run, size=12)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def add_figure(doc, path: Path, caption: str, width_in=6.0):
    if not path.exists():
        add_para(doc, f"[Khong tim thay hinh: {path.name}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(optimized(path)), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run_font(run, size=11, italic=True)
    doc.add_paragraph()


def main():
    year = 2018
    fig_dir = cfg.FIGURES_DIR

    print("Doc lai du lieu that de lay so lieu dua vao bao cao...")
    meta = validation.get_raster_metadata(cfg.raw_image_path(year))
    band_stats = validation.compute_band_statistics(cfg.raw_image_path(year), cfg.BAND_NAMES)
    samples, counts = preprocessing.sample_valid_pixels(year, n_samples=cfg.N_SAMPLES_PER_YEAR)
    df = features.build_feature_dataframe(year, samples)
    corr = df[cfg.FEATURE_NAMES].corr()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(13)
    normal.font.color.rgb = BLACK

    # --- Trang tieu de --------------------------------------------------
    add_para(
        doc,
        "PHÂN TÍCH DỮ LIỆU ẢNH VỆ TINH SENTINEL-2 NĂM 2018\n"
        "KHU VỰC SÂN BAY LONG THÀNH VÀ VÙNG PHỤ CẬN",
        size=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "(Dữ liệu nền / baseline trước khi khởi công dự án)",
        size=13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(doc, f"Ngày lập báo cáo: {date.today().strftime('%d/%m/%Y')}", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()

    # --- 1. Giới thiệu ----------------------------------------------------
    add_heading(doc, "1. Giới thiệu", 1)
    add_para(
        doc,
        "Báo cáo này phân tích ảnh vệ tinh Sentinel-2 chụp khu vực Sân bay Long Thành và vùng "
        "phụ cận (Đồng Nai) vào năm 2018 — thời điểm trước khi dự án sân bay khởi công xây dựng. "
        "Đây là dữ liệu nền (baseline), dùng để làm mốc so sánh cho các năm tiếp theo khi phân "
        "tích tốc độ mở rộng khu vực xây dựng/bê tông hóa quanh sân bay."
    )
    add_para(
        doc,
        "Phần phân tích dưới đây thuộc Phase 1 của dự án (tiền xử lý dữ liệu và khám phá dữ liệu "
        "- EDA), chưa thực hiện phân loại lớp phủ hay tính tốc độ bê tông hóa. Toàn bộ số liệu "
        "trong báo cáo được lấy trực tiếp từ ảnh vệ tinh thật, không có số liệu giả định."
    )

    # --- 2. Nguồn dữ liệu ---------------------------------------------------
    add_heading(doc, "2. Nguồn dữ liệu", 1)
    add_para(
        doc,
        "Ảnh được lấy từ kho dữ liệu công khai Sentinel-2 L2A trên AWS (thông qua Earth Search "
        "STAC API của Element84), không cần tài khoản hay trả phí. Ảnh đã qua hiệu chỉnh khí "
        "quyển (surface reflectance, mức xử lý L2A)."
    )
    add_table(
        doc,
        ["Thuộc tính", "Giá trị"],
        [
            ["Mã cảnh (scene id)", "S2A_48PYS_20181129_0_L2A"],
            ["Ngày chụp", "29/11/2018"],
            ["Vệ tinh", "Sentinel-2A"],
            ["Tile MGRS", "48PYS"],
            ["Hệ tọa độ (CRS)", meta["crs"]],
            ["Độ phân giải không gian", f"{meta['resolution_x']:.1f} m x {meta['resolution_y']:.1f} m"],
            ["Kích thước ảnh (AOI)", f"{meta['width']} x {meta['height']} pixel"],
            ["Diện tích vùng nghiên cứu", "khoảng 30 km x 21 km quanh sân bay Long Thành"],
            ["Số kênh phổ (band)", f"{meta['band_count']} ({', '.join(cfg.BAND_NAMES)})"],
            ["Kiểu dữ liệu", meta["dtype"]],
        ],
    )

    # --- 3. Tiền xử lý -------------------------------------------------------
    add_heading(doc, "3. Tiền xử lý dữ liệu", 1)
    add_para(
        doc,
        "Ảnh gốc được cắt (window read) trực tiếp từ ảnh Sentinel-2 gốc theo ranh giới vùng "
        "nghiên cứu (AOI), gồm 6 kênh phổ B2, B3, B4, B8, B11, B12. Hai kênh B11 và B12 vốn có "
        "độ phân giải gốc 20m, được resample lên 10m để đồng bộ lưới pixel với các kênh còn lại."
    )
    add_para(
        doc,
        "Về mây: ảnh Sentinel-2 L2A đi kèm lớp phân loại cảnh (Scene Classification Layer - SCL) "
        "do ESA tạo sẵn, dùng thuật toán Sen2Cor. Trong báo cáo này, mặt nạ mây được lấy trực "
        "tiếp từ SCL (các lớp mây, bóng mây, cirrus mỏng), không phải mặt nạ tự tạo."
    )
    add_para(
        doc,
        "Sau khi đọc dữ liệu thô, từng pixel được phân loại thành 1 trong các nhóm: hợp lệ "
        "(valid), bị mây che, nằm ngoài AOI, hoặc có giá trị phản xạ vượt ngưỡng hợp lệ "
        "(0-10000). Không có pixel nào bị loại bỏ âm thầm — mọi pixel không hợp lệ đều được "
        "đếm và báo cáo lại như bảng dưới đây."
    )

    # --- 4. Chất lượng dữ liệu -------------------------------------------------
    add_heading(doc, "4. Đánh giá chất lượng dữ liệu", 1)
    total = counts["total"]
    add_table(
        doc,
        ["Loại pixel", "Số lượng", "Tỷ lệ (%)"],
        [
            ["Tổng số pixel", f"{total:,}", "100.00"],
            ["Hợp lệ (valid)", f"{counts['valid']:,}", f"{100*counts['valid']/total:.2f}"],
            ["Mây / bóng mây / cirrus (từ SCL)", f"{counts['cloud']:,}", f"{100*counts['cloud']/total:.2f}"],
            ["Nằm ngoài AOI", f"{counts['outside_aoi']:,}", f"{100*counts['outside_aoi']/total:.2f}"],
            ["Vượt ngưỡng phản xạ hợp lệ", f"{counts['invalid_range']:,}", f"{100*counts['invalid_range']/total:.2f}"],
            ["Thiếu dữ liệu (NoData)", f"{counts['missing']:,}", f"{100*counts['missing']/total:.2f}"],
        ],
    )
    add_para(
        doc,
        f"Không có pixel NoData (0.00%) — ảnh phủ kín toàn bộ vùng nghiên cứu. Tỷ lệ mây thực "
        f"tế trong AOI là {100*counts['cloud']/total:.2f}%, cao hơn con số mây trung bình mà "
        f"metadata STAC báo cho toàn cảnh gốc (2.2%), vì vùng AOI nằm ở phần có mây cục bộ nhiều "
        f"hơn phần còn lại của cảnh. Số pixel hợp lệ còn lại để phân tích là "
        f"{100*counts['valid']/total:.2f}%, đủ lớn để thống kê đáng tin cậy."
    )
    add_para(
        doc,
        f"Có {counts['invalid_range']:,} pixel ({100*counts['invalid_range']/total:.2f}%) có giá "
        "trị phản xạ vượt ngưỡng 0-10000, nhiều khả năng là viền mây hoặc hiện tượng phản xạ "
        "chói (sun glint) mà lớp SCL chưa bắt hết. Các pixel này được giữ lại trong dữ liệu gốc "
        "và gắn cờ, không bị xóa, để tránh làm sai lệch thống kê một cách âm thầm."
    )

    # --- 5. Thống kê phổ --------------------------------------------------
    add_heading(doc, "5. Thống kê giá trị phản xạ theo từng kênh", 1)
    add_para(
        doc,
        "Bảng dưới đây tính trên mẫu 20.000 pixel hợp lệ được lấy ngẫu nhiên (có seed cố định "
        "để tái lập được), giá trị phản xạ đã quy về thang 0-1."
    )
    rows = []
    for b in cfg.BAND_NAMES:
        s = df[b]
        rows.append([b, f"{s.mean():.3f}", f"{s.std():.3f}", f"{s.min():.3f}", f"{s.median():.3f}", f"{s.max():.3f}"])
    add_table(doc, ["Kênh", "Trung bình", "Độ lệch chuẩn", "Min", "Trung vị", "Max"], rows)
    add_para(
        doc,
        "Kênh cận hồng ngoại B8 (NIR) có giá trị trung bình cao hơn hẳn các kênh khả kiến "
        "(B2, B3, B4) — đây là đặc điểm điển hình của vùng có thảm thực vật chiếm ưu thế, phù "
        "hợp với bối cảnh khu vực trước khi giải phóng mặt bằng xây sân bay."
    )

    add_figure(doc, fig_dir / f"band_distribution_{year}.png", "Hình 1. Phân bố giá trị phản xạ của 6 kênh phổ B2-B12, năm 2018.")

    # --- 6. Chỉ số phổ -----------------------------------------------------
    add_heading(doc, "6. Các chỉ số phổ NDVI, NDBI, NDWI", 1)
    add_para(doc, "Ba chỉ số được tính từ các kênh phổ theo công thức chuẩn:")
    add_para(doc, "NDVI = (NIR - Red) / (NIR + Red)   —  chỉ số thực vật")
    add_para(doc, "NDBI = (SWIR1 - NIR) / (SWIR1 + NIR)   —  chỉ số khu vực xây dựng")
    add_para(doc, "NDWI = (Green - NIR) / (Green + NIR)   —  chỉ số mặt nước")
    add_para(
        doc,
        "Công thức có xử lý chia cho 0 (trả về NaN thay vì lỗi hoặc giá trị vô hạn), tránh làm "
        "sai lệch thống kê khi mẫu số bằng 0."
    )
    rows = []
    for idx_name in cfg.INDEX_NAMES:
        s = df[idx_name]
        rows.append([idx_name, f"{s.mean():.3f}", f"{s.std():.3f}", f"{s.min():.3f}", f"{s.median():.3f}", f"{s.max():.3f}"])
    add_table(doc, ["Chỉ số", "Trung bình", "Độ lệch chuẩn", "Min", "Trung vị", "Max"], rows)
    add_para(
        doc,
        f"NDVI có trung vị khá cao ({df['NDVI'].median():.2f}), cho thấy thực vật chiếm phần lớn "
        f"diện tích năm 2018. NDBI trung vị âm ({df['NDBI'].median():.2f}) cho thấy khu vực xây "
        f"dựng lúc này còn rất ít — hợp lý vì đây là ảnh trước khi dự án khởi công. NDWI trung vị "
        f"âm khá sâu ({df['NDWI'].median():.2f}) cho thấy diện tích mặt nước trong AOI không lớn."
    )
    add_para(
        doc,
        "Lưu ý: đây chỉ là quan sát mô tả trên dữ liệu phổ thô, chưa phải kết quả phân loại lớp "
        "phủ, nên không dùng để kết luận chính xác diện tích từng loại đất."
    )
    add_figure(doc, fig_dir / f"index_distribution_{year}.png", "Hình 2. Phân bố NDVI, NDBI, NDWI, năm 2018.")
    add_figure(doc, fig_dir / f"feature_boxplots_{year}.png", "Hình 3. Boxplot các kênh phổ và chỉ số, dùng để xem outlier và độ phân tán.")

    # --- 7. Tương quan --------------------------------------------------
    add_heading(doc, "7. Phân tích tương quan giữa các đặc trưng", 1)
    add_para(
        doc,
        "Ma trận tương quan cho thấy một số cặp đặc trưng có quan hệ khá chặt, có thể trùng lặp "
        "thông tin khi đưa vào mô hình phân loại ở Phase 2:"
    )
    import numpy as np
    tri = corr.where(np.triu(np.ones(corr.shape), k=1).astype(bool)).stack().abs().sort_values(ascending=False)
    rows = [[f"{a} - {b}", f"{corr.loc[a,b]:.2f}"] for (a, b) in tri.index[:6]]
    add_table(doc, ["Cặp đặc trưng", "Hệ số tương quan (r)"], rows)
    add_para(
        doc,
        "Ba kênh khả kiến B2, B3, B4 tương quan rất cao với nhau (r trên 0.9), tương tự với cặp "
        "SWIR B11-B12. Đáng chú ý là NDVI và NDWI tương quan nghịch gần như tuyệt đối "
        f"(r = {corr.loc['NDVI','NDWI']:.2f}) trong ảnh này — có thể vì diện tích mặt nước trong "
        "AOI khá nhỏ nên hai chỉ số gần như đối lập nhau. Những cặp tương quan cao này nên được "
        "cân nhắc khi chọn đặc trưng đưa vào Random Forest ở phase sau, tránh dư thừa thông tin."
    )
    add_figure(doc, fig_dir / f"correlation_heatmap_{year}.png", "Hình 4. Ma trận tương quan giữa các kênh phổ và chỉ số, năm 2018.")

    # --- 8. Trực quan không gian --------------------------------------------
    add_heading(doc, "8. Trực quan hóa không gian", 1)
    add_para(doc, "Ảnh tổ hợp màu tự nhiên (RGB = B4-B3-B2) của khu vực nghiên cứu năm 2018:")
    add_figure(doc, fig_dir / f"rgb_{year}.png", "Hình 5. Ảnh RGB tổ hợp (B4-B3-B2), năm 2018.")

    add_para(doc, "Mặt nạ mây thực tế lấy từ lớp SCL của Sentinel-2 (vùng màu đen là mây/bóng mây/cirrus):")
    add_figure(doc, fig_dir / f"cloudmask_{year}.png", "Hình 6. Mặt nạ mây (SCL), năm 2018.")

    add_para(doc, "Bản đồ không gian của 3 chỉ số phổ:")
    add_figure(doc, fig_dir / f"ndvi_map_{year}.png", "Hình 7. Bản đồ NDVI, năm 2018.")
    add_figure(doc, fig_dir / f"ndbi_map_{year}.png", "Hình 8. Bản đồ NDBI, năm 2018.")
    add_figure(doc, fig_dir / f"ndwi_map_{year}.png", "Hình 9. Bản đồ NDWI, năm 2018.")
    add_para(
        doc,
        "Trên bản đồ NDVI, các vùng màu xanh đậm (giá trị cao) chiếm phần lớn diện tích, khớp "
        "với nhận định thực vật chiếm ưu thế ở phần thống kê trên. Bản đồ NDBI không có vùng "
        "giá trị cao rõ rệt lan rộng, phù hợp với việc khu vực chưa xây dựng nhiều vào năm 2018."
    )

    # --- 9. Hạn chế ------------------------------------------------------
    add_heading(doc, "9. Một số hạn chế của phân tích này", 1)
    add_para(doc, "- Chỉ dùng 1 cảnh ảnh duy nhất cho năm 2018 (ngày 29/11/2018), không phải ảnh ghép nhiều thời điểm, nên kết quả có thể bị ảnh hưởng bởi điều kiện khí quyển/mây của đúng ngày chụp đó.")
    add_para(doc, "- Tỷ lệ 1.41% pixel \"ngoài AOI\" là do AOI hình chữ nhật (tọa độ WGS84) không khớp hoàn toàn với lưới pixel UTM, không phải lỗi dữ liệu.")
    add_para(doc, "- 0.25% pixel vượt ngưỡng phản xạ hợp lệ vẫn chưa được xử lý dứt điểm, cần quyết định ở Phase 2 là loại bỏ hay giữ lại khi huấn luyện mô hình.")
    add_para(doc, "- Đây mới là phân tích mô tả (EDA), chưa có nhãn lớp phủ (land-cover label) nên chưa thể nói chính xác bao nhiêu % diện tích là đất trống, nước, thực vật hay xây dựng.")

    # --- 10. Kết luận -------------------------------------------------------
    add_heading(doc, "10. Kết luận", 1)
    add_para(
        doc,
        "Ảnh Sentinel-2 năm 2018 tải về đủ điều kiện để làm dữ liệu nền cho dự án: không có "
        "NoData, tỷ lệ pixel hợp lệ cao (86.9%), mây ở mức chấp nhận được (khoảng 12%) và có "
        "mặt nạ mây thật đi kèm. Các chỉ số phổ cho thấy bức tranh hợp lý về mặt vật lý: thực "
        "vật chiếm ưu thế, khu vực xây dựng còn thấp — đúng như kỳ vọng trước khi dự án sân bay "
        "khởi công."
    )
    add_para(
        doc,
        "Dữ liệu năm 2018 này sẽ được dùng làm mốc so sánh khi phân tích các năm 2020, 2022, "
        "2024, 2026 (đã tải sẵn nhưng chưa phân tích) để quan sát xu hướng thay đổi chỉ số phổ "
        "theo thời gian ở các phần sau của Phase 1, trước khi chuyển sang huấn luyện mô hình "
        "Random Forest ở Phase 2."
    )

    doc.save(OUT_PATH)
    print(f"[done] Da ghi {OUT_PATH}")


if __name__ == "__main__":
    main()
